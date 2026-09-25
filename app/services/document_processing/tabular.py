"""Safe tabular evidence readers shared by extraction and audit validators."""
import csv
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from app.services.audit_engine.header_matching import match_fields_exclusive


MAX_CONSECUTIVE_EMPTY_ROWS = 200
MAX_LEADING_EMPTY_ROWS = 1_000


def cell_text(value: object) -> str:
    """Render a cell without treating valid zero values as blank."""
    return '' if value is None else str(value)


def _meaningful_rows(rows):
    """Retain populated rows without walking arbitrarily formatted tails.

    Excel templates often apply formatting through row 1,048,576 even though
    their evidence table ends after a few records.  A 200-row empty boundary
    preserves normal table spacing while preventing an audit scan from being
    held hostage by formatting-only cells.
    """
    result, empty_rows = [], 0
    for row in rows:
        values = list(row)
        if any(cell_text(value).strip() for value in values):
            result.append(values)
            empty_rows = 0
        else:
            empty_rows += 1
            limit = MAX_CONSECUTIVE_EMPTY_ROWS if result else MAX_LEADING_EMPTY_ROWS
            if empty_rows >= limit:
                break
    return result


@dataclass(frozen=True)
class TabularSelection:
    """The most likely evidence table within a spreadsheet workbook."""
    sheet_name: str
    header_row_index: int
    rows: list[list]
    field_indices: dict[str, int | None]


def read_tabular_sheets(path: str, *, data_only: bool = True) -> list[tuple[str, list[list]]]:
    """Read every worksheet while preserving its name and row order.

    Customer workbooks commonly keep revision history on the first worksheet
    and the governed register on a later one.  Returning all sheets prevents
    a version-history tab from being assessed as the actual evidence.
    """
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == '.csv':
        with source.open('r', encoding='utf-8-sig', errors='replace', newline='') as stream:
            return [('CSV', [row for row in csv.reader(stream)])]
    if suffix == '.xls':
        import xlrd
        workbook = xlrd.open_workbook(source)
        return [(sheet.name, _meaningful_rows(sheet.row_values(index) for index in range(sheet.nrows)))
                for sheet in workbook.sheets()]
    workbook = load_workbook(source, read_only=True, data_only=data_only, keep_links=False)
    try:
        return [(sheet.title, _meaningful_rows(sheet.iter_rows(values_only=True)))
                for sheet in workbook.worksheets]
    finally:
        workbook.close()


def read_tabular_rows(path: str) -> list[list]:
    """Return the first sheet for backwards-compatible callers.

    Evidence validators must use :func:`select_tabular_table` instead; it
    evaluates every worksheet and finds the actual header row.
    """
    sheets = read_tabular_sheets(path)
    return sheets[0][1] if sheets else []


def select_tabular_tables(path: str, fields: list[dict], required_fields: tuple[str, ...] = (),
                          header_search_rows: int = 100, *, data_only: bool = True) -> list[TabularSelection]:
    """Return the best governed table from every relevant worksheet.

    A single workbook can legitimately hold a record register, supporting
    approval data, and a calculation table on separate sheets.  Callers that
    validate every master control must examine each relevant table, not only
    the workbook's globally best header match.
    """
    selected: list[tuple[tuple[int, int, int], TabularSelection]] = []
    for sheet_name, sheet_rows in read_tabular_sheets(path, data_only=data_only):
        sheet_best: TabularSelection | None = None
        sheet_best_score = (-1, -1, -1)
        for row_index, headers in enumerate(sheet_rows[:header_search_rows]):
            # A record can contain free-text values that look like field names
            # in an adjacent summary block. Dates and numeric cell values are
            # a reliable indication that this is a data row, never the table
            # header. Skipping it prevents the first incident/defect record
            # from becoming the schema for an entire workbook.
            if any(
                value is not None and not isinstance(value, str)
                for value in headers
            ):
                continue
            indices = match_fields_exclusive(headers, fields)
            required_matches = sum(indices.get(key) is not None for key in required_fields)
            matches = sum(index is not None for index in indices.values())
            populated = sum(bool(cell_text(value).strip()) for value in headers)
            score = (required_matches, matches, populated)
            if score > sheet_best_score:
                sheet_best_score = score
                sheet_best = TabularSelection(sheet_name, row_index, sheet_rows[row_index:], indices)
        # A zero-match row is not a governed table for this document schema.
        # Excluding it avoids treating a revision-history sheet as evidence.
        if sheet_best is not None and sheet_best_score[1] > 0:
            selected.append((sheet_best_score, sheet_best))
    return [selection for _, selection in sorted(selected, key=lambda item: item[0], reverse=True)]


def select_tabular_table(path: str, fields: list[dict], required_fields: tuple[str, ...] = (),
                         header_search_rows: int = 100, *, data_only: bool = True) -> TabularSelection:
    """Select the strongest governed table for backwards-compatible callers."""
    selections = select_tabular_tables(
        path, fields, required_fields, header_search_rows, data_only=data_only
    )
    if not selections:
        return TabularSelection('', 0, [], {field['key']: None for field in fields})
    return selections[0]
