"""Validate, version, and safely extend the governed CMMI rules catalogue."""
from __future__ import annotations

from datetime import date
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from sqlalchemy import select, update
from sqlalchemy.orm import Session

from app.db.models import ChecklistVersion, CmmiRule, DocumentTypeRule, PracticeArea
from app.services.document_processing.file_security import validate_zip

PA_ALIASES = {'VER': 'VV', 'VAL': 'VV', 'MC': 'MPM', 'SAM': 'SUP', 'SSS': 'SUP', 'PPL': 'OT'}
MAX_RULE_WORKBOOK_BYTES = 50 * 1024 * 1024
MAX_RULE_WORKBOOK_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_DOCUMENT_TYPE_LENGTH = 500
MAX_CATALOGUE_TEXT_LENGTH = 10_000
MAX_CATALOGUE_SHEET_COLUMNS = 200
DOCUMENT_ADDITION_COLUMNS = (
    'Document / Evidence',
    'Related Practice Area',
    'What to Check / Key Points',
    'Expected Evidence',
)
SAP_MASTER_COLUMNS = (
    'Document / Artefact', 'CMMI Practice Area', 'Rule ID',
    'Audit Check / Required Evidence (Field-Level)', 'Mandatory', 'Condition',
    'What Tool Should Detect', 'AFR Finding if Missing', 'Recommended Action',
)
# The supplied SAP workbook deliberately combines Change Request and FSD in
# its 14-document index, while its detailed master gives each artefact its
# own section.  These are the authoritative links from master sections to the
# corresponding governed evidence type.
SAP_MASTER_DOCUMENT_MAP = {
    'changerequest': 'changelogfsd',
    'changerequestlogfsd': 'changelogfsd',
    'changelogfsd': 'changelogfsd',
    'impactanalysistsd': 'impactanalysisandtsd',
    'estimationsheet': 'estimationsheet',
    'wbs': 'wbs',
    'trackingofactualeffortagainstplanned': 'trackingofactualeffortsagainstplanned',
    'unittestcasesandresults': 'unittestcasesresults',
    'codereviewrecordsanddefects': 'codereviewrecordsdefects',
    'testcasesandresults': 'testcasesresults',
    'testingdefects': 'testingdefects',
    'transportrequestwithreleasenote': 'transportrequestreleasenote',
    'incidentlog2months': 'incidentlog2months',
    'slatargetsvsachievement': 'slatargetvsachievementreport',
    'camsheetcorrectiveactionmonitoring': 'camsheet',
    'rcaforp1p2': 'rcaforp1p2',
}
_APPROVED_DOCUMENT_CATALOGUE = (
    Path(__file__).resolve().parents[1] / 'db' / 'seed_data' / 'document_types.json'
)


def _text(value) -> str:
    return re.sub(r'\s+', ' ', str(value or '')).strip()


def _norm(value) -> str:
    return re.sub(r'[^a-z0-9]+', '', _text(value).lower())


def _approved_documents() -> list[dict]:
    """Return the governed 14-entry evidence catalogue.

    The master CMMI workbook owns detailed CMMI questions.  It does *not*
    own the evidence-document catalogue: that catalogue is intentionally a
    small, approved set so an uploaded master workbook cannot silently widen
    the scan or AFR scope.
    """
    try:
        source = json.loads(_APPROVED_DOCUMENT_CATALOGUE.read_text(encoding='utf-8'))
    except (OSError, json.JSONDecodeError) as exc:
        raise RuntimeError('The approved document catalogue seed is unavailable or invalid.') from exc
    documents = []
    seen = set()
    for item in source:
        name = _text(item.get('document_type'))
        key = _norm(name)
        if not key or key in seen:
            raise RuntimeError('The approved document catalogue contains an invalid or duplicate document type.')
        if not item.get('practice_areas') or not item.get('keywords'):
            raise RuntimeError(f'The approved document catalogue entry "{name}" is incomplete.')
        seen.add(key)
        documents.append({
            'document_type': name,
            'practice_areas': list(item['practice_areas']),
            'keywords': list(item['keywords']),
            'aliases': list(item.get('aliases') or []),
            'expected_evidence': item.get('expected_evidence') or '',
            'primary_purpose': item.get('primary_purpose'),
            'include_in_afr': True,
        })
    if len(documents) != 14:
        raise RuntimeError('The approved document catalogue must contain exactly 14 document types.')
    return documents


def _rows(sheet, maximum: int = 10_000) -> list[list]:
    # Check declared dimensions before iteration. A tiny malicious worksheet
    # can claim Excel's full grid and make ``iter_rows`` materialize millions
    # of empty cells despite passing the compressed/uncompressed byte limits.
    if sheet.max_row > maximum:
        raise ValueError(f'Sheet {sheet.title} exceeds the {maximum}-row import limit.')
    if sheet.max_column > MAX_CATALOGUE_SHEET_COLUMNS:
        raise ValueError(
            f'Sheet {sheet.title} exceeds the {MAX_CATALOGUE_SHEET_COLUMNS}-column import limit.'
        )
    rows = []
    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if index > maximum:
            raise ValueError(f'Sheet {sheet.title} exceeds the {maximum}-row import limit.')
        rows.append(list(row))
    return rows


def _header(rows: list[list], required: str) -> tuple[int, list]:
    wanted = _norm(required)
    for index, row in enumerate(rows):
        if any(_norm(cell) == wanted for cell in row):
            return index, row
    raise ValueError(f'Missing required header "{required}".')


def _column(header: list, name: str) -> int:
    wanted = _norm(name)
    matches = [index for index, value in enumerate(header) if _norm(value) == wanted]
    if not matches:
        raise ValueError(f'Missing required column "{name}".')
    if len(matches) > 1:
        raise ValueError(f'Duplicate required column "{name}".')
    return matches[0]


def _split_keywords(value) -> list[str]:
    return list(dict.fromkeys(item for item in (_text(part) for part in re.split(r'[;,\n]', str(value or ''))) if len(item) > 1))


def _split_expected_evidence(value) -> list[str]:
    return list(dict.fromkeys(item for item in (_text(part) for part in re.split(r'[;\n]', str(value or ''))) if item))


def _merge_text_values(*groups: list[str]) -> list[str]:
    """Return ordered, case-insensitive unique text while preserving existing data."""
    values, seen = [], set()
    for group in groups:
        for value in group:
            clean = _text(value)
            key = _norm(clean)
            if clean and key not in seen:
                values.append(clean)
                seen.add(key)
    return values


def _canonical_practice_area(code: str, known: set[str]) -> str:
    """Prefer an exact governed code and use legacy aliases only as fallback."""
    return code if code in known else PA_ALIASES.get(code, code)


def _addition_practice_areas(value: str, known: set[str], location: str) -> list[str]:
    """Read either a code list or a human-readable PA label safely."""
    source = value.upper().strip()
    raw_codes = re.findall(r'\b[A-Z]{2,5}\b', source)
    mapped_codes = [_canonical_practice_area(code, known) for code in raw_codes]
    recognized = list(dict.fromkeys(code for code in mapped_codes if code in known))
    # When the whole cell is a compact list such as RSK/EST, every token is
    # intended to be a code and typos must fail. In a label such as
    # "Risk Management (RSK)", ordinary words must not be mistaken for codes.
    is_code_list = bool(re.fullmatch(
        r'\s*[A-Z]{2,5}(?:\s*[/,;&+]\s*[A-Z]{2,5})*\s*', source
    ))
    invalid = sorted(set(code for code in mapped_codes if code not in known))
    if is_code_list and invalid:
        raise ValueError(f'{location} has unknown practice area(s): {", ".join(invalid)}.')
    if not recognized:
        raise ValueError(f'{location} must include at least one governed practice area code.')
    return recognized


def _load_safe_workbook(content: bytes, error_message: str):
    if not content or len(content) > MAX_RULE_WORKBOOK_BYTES:
        raise ValueError('Workbook is empty or exceeds the 50 MB service limit.')
    try:
        validate_zip(
            BytesIO(content),
            max_members=2_000,
            max_uncompressed=MAX_RULE_WORKBOOK_UNCOMPRESSED_BYTES,
            max_ratio=100,
        )
        return load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValueError(error_message) from exc


def _parse_rules(workbook) -> list[dict]:
    if 'MASTER' not in workbook.sheetnames:
        raise ValueError('Workbook must contain a MASTER sheet.')
    rows = _rows(workbook['MASTER'])
    start, header = _header(rows, 'Rule ID')
    level_at = _column(header, 'Level')
    rule_at = _column(header, 'Rule ID')
    check_at = _column(header, 'Audit Check (What the auditor looks for and verifies)')
    gap_at = _column(header, 'If NO: Finding and Implication')
    rules, seen = [], set()
    for row in rows[start + 1:]:
        rule_id = _text(row[rule_at] if rule_at < len(row) else '').upper()
        if not re.fullmatch(r'[A-Z0-9]+-[A-Z0-9]+', rule_id, re.I):
            continue
        duplicate_key = rule_id.casefold()
        if duplicate_key in seen:
            raise ValueError(f'Duplicate Rule ID: {rule_id}.')
        seen.add(duplicate_key)
        audit_check = _text(row[check_at] if check_at < len(row) else '')
        gap_text = _text(row[gap_at] if gap_at < len(row) else '')
        if not audit_check or not gap_text:
            raise ValueError(f'Rule {rule_id} is missing its audit check or finding text.')
        level = _text(row[level_at] if level_at < len(row) else '')
        if level not in {'L1-Gate', 'L2-Block', 'L3-Check', 'L4-Probe'}:
            raise ValueError(f'Rule {rule_id} has an invalid level: {level or "blank"}.')
        rules.append({'rule_id': rule_id, 'practice_area': rule_id.split('-')[0].upper(), 'level': level,
                      'is_gate': level.lower() == 'l1-gate', 'is_check': level.lower() == 'l3-check',
                      'is_probe': level.lower() == 'l4-probe', 'audit_check': audit_check, 'gap_text': gap_text})
    if not rules:
        raise ValueError('MASTER contains no valid audit rules.')
    return rules


def _document_aliases(name: str) -> list[str]:
    return _merge_text_values([name], [_text(part) for part in name.split('/')])


def _parse_document_additions(workbook, known: set[str]) -> list[dict]:
    """Parse a standalone evidence-catalogue sheet without requiring MASTER.

    This path is intentionally additive: it accepts only the document columns
    and never creates or modifies CMMI rule records.
    """
    for sheet in workbook.worksheets:
        rows = _rows(sheet)
        try:
            start, header = _header(rows, 'Document / Evidence')
        except ValueError:
            continue
        indexes = {name: _column(header, name) for name in DOCUMENT_ADDITION_COLUMNS}
        additions: dict[str, dict] = {}
        for row_number, row in enumerate(rows[start + 1:], start=start + 2):
            values = {name: _text(row[index] if index < len(row) else '') for name, index in indexes.items()}
            if not any(values.values()):
                continue
            if any(value.lstrip().startswith('=') for value in values.values()):
                raise ValueError(f'{sheet.title} row {row_number} contains a formula. Catalogue values must be plain text.')
            if any('\ufffd' in value for value in values.values()):
                raise ValueError(f'{sheet.title} row {row_number} contains an invalid replacement character. Correct the source text before importing.')
            name = values['Document / Evidence'].rstrip('/')
            if not name:
                raise ValueError(f'{sheet.title} row {row_number} is missing Document / Evidence.')
            if len(name) > MAX_DOCUMENT_TYPE_LENGTH:
                raise ValueError(f'{sheet.title} row {row_number} exceeds the {MAX_DOCUMENT_TYPE_LENGTH}-character document name limit.')
            if any(len(value) > MAX_CATALOGUE_TEXT_LENGTH for value in values.values()):
                raise ValueError(f'{sheet.title} row {row_number} contains a value longer than {MAX_CATALOGUE_TEXT_LENGTH} characters.')
            key = _norm(name)
            if not key:
                raise ValueError(f'{sheet.title} row {row_number} has a document name without searchable letters or numbers.')
            practice_areas = _addition_practice_areas(
                values['Related Practice Area'], known, f'{sheet.title} row {row_number}'
            )
            keywords = _split_keywords(values['What to Check / Key Points'])
            expected_evidence = values['Expected Evidence']
            if not keywords:
                raise ValueError(f'{sheet.title} row {row_number} is missing What to Check / Key Points.')
            if not expected_evidence:
                raise ValueError(f'{sheet.title} row {row_number} is missing Expected Evidence.')
            item = additions.setdefault(key, {
                'document_type': name,
                'practice_areas': [],
                'keywords': [],
                'aliases': [],
                'expected_evidence': [],
            })
            item['practice_areas'] = _merge_text_values(item['practice_areas'], practice_areas)
            item['keywords'] = _merge_text_values(item['keywords'], keywords)
            item['aliases'] = _merge_text_values(item['aliases'], _document_aliases(name))
            item['expected_evidence'] = _merge_text_values(item['expected_evidence'], _split_expected_evidence(expected_evidence))
        if not additions:
            raise ValueError(f'{sheet.title} contains no document/evidence rows.')
        return list(additions.values())
    raise ValueError('Workbook must contain a sheet with the document/evidence catalogue columns.')


def _parse_sap_audit_checklist(workbook, known: set[str]) -> tuple[list[dict], list[dict]]:
    """Read the final SAP workbook as 14 evidence types and field controls.

    The index supplies the document names, keyword schema, and expected
    evidence.  The detailed master is authoritative for the document mapping,
    field-level rule, applicability condition, detection instruction, finding,
    and remediation text.
    """
    if 'Master Checklist' not in workbook.sheetnames:
        raise ValueError('SAP checklist must contain a "Master Checklist" sheet.')
    additions = _parse_document_additions(workbook, known)
    documents_by_key = {_norm(item['document_type']): item for item in additions}
    if len(documents_by_key) != 14:
        raise ValueError('SAP checklist must contain exactly 14 distinct document/evidence entries.')
    master_rows = _rows(workbook['Master Checklist'])
    start, header = _header(master_rows, 'Document / Artefact')
    indexes = {name: _column(header, name) for name in SAP_MASTER_COLUMNS}
    controls: list[dict] = []
    seen_rule_ids: set[str] = set()
    seen_document_keys: set[str] = set()
    master_practice_areas: dict[str, set[str]] = {}
    master_aliases: dict[str, list[str]] = {}
    for row_number, row in enumerate(master_rows[start + 1:], start=start + 2):
        values = {name: _text(row[index] if index < len(row) else '') for name, index in indexes.items()}
        if not any(values.values()):
            continue
        if any(value.lstrip().startswith('=') for value in values.values()):
            raise ValueError(f'Master Checklist row {row_number} contains a formula. Controls must be plain text.')
        artifact = values['Document / Artefact']
        document_key = SAP_MASTER_DOCUMENT_MAP.get(_norm(artifact))
        if not document_key or document_key not in documents_by_key:
            raise ValueError(f'Master Checklist row {row_number} has an unmapped document/artefact: {artifact or "blank"}.')
        practice_area = values['CMMI Practice Area'].upper()
        if practice_area not in known:
            raise ValueError(f'Master Checklist row {row_number} has an unknown practice area: {practice_area or "blank"}.')
        rule_id = values['Rule ID'].upper()
        if not re.fullmatch(r'[A-Z0-9]+-[A-Z0-9]+', rule_id):
            raise ValueError(f'Master Checklist row {row_number} has an invalid Rule ID: {rule_id or "blank"}.')
        if rule_id in seen_rule_ids:
            raise ValueError(f'Master Checklist has a duplicate Rule ID: {rule_id}.')
        audit_check = values['Audit Check / Required Evidence (Field-Level)']
        gap_text = values['AFR Finding if Missing']
        detection = values['What Tool Should Detect']
        recommendation = values['Recommended Action']
        if not all((audit_check, gap_text, detection, recommendation)):
            raise ValueError(f'Master Checklist row {row_number} is missing a required control, detection, finding, or action value.')
        mandatory = values['Mandatory'].casefold() in {'yes', 'y', 'true', 'mandatory'}
        if not mandatory:
            raise ValueError(f'Master Checklist row {row_number} must mark the control as mandatory.')
        seen_rule_ids.add(rule_id)
        seen_document_keys.add(document_key)
        master_practice_areas.setdefault(document_key, set()).add(practice_area)
        master_aliases.setdefault(document_key, []).append(artifact)
        controls.append({
            'rule_id': rule_id, 'practice_area': practice_area, 'level': 'L3-Check',
            'document_key': document_key, 'audit_check': audit_check, 'gap_text': gap_text,
            'condition': values['Condition'] or 'Always', 'detection': detection,
            'recommendation': recommendation, 'is_mandatory': mandatory,
        })
    if len(controls) != 210:
        raise ValueError(f'SAP Master Checklist must contain 210 controls; found {len(controls)}.')
    if len(seen_document_keys) != 14:
        raise ValueError('SAP Master Checklist must map controls to all 14 approved document types.')
    documents = []
    for key, item in documents_by_key.items():
        practice_areas = sorted(master_practice_areas.get(key, set()))
        if not practice_areas:
            raise ValueError(f'Sheet1 document "{item["document_type"]}" has no detailed Master Checklist controls.')
        documents.append({
            'document_type': item['document_type'], 'practice_areas': practice_areas,
            'keywords': item['keywords'],
            'aliases': _merge_text_values(item['aliases'], master_aliases.get(key, [])),
            'expected_evidence': '; '.join(item['expected_evidence']),
            'primary_purpose': None, 'include_in_afr': True,
        })
    return documents, controls


def _rule_signature(rule: dict) -> tuple:
    # A field-control's detection, condition, and action are executable audit
    # behaviour, not display-only wording.  Any change must create a new
    # governed version rather than being discarded as a no-op.
    return (
        rule['practice_area'], rule['level'], rule['audit_check'], rule['gap_text'],
        rule.get('document_key'), rule.get('condition'), rule.get('detection'),
        rule.get('recommendation'), rule.get('is_mandatory', True),
    )


def _document_signature(document: dict) -> tuple:
    return (
        tuple(document.get('practice_areas') or []),
        tuple(document.get('keywords') or []),
        tuple(document.get('aliases') or []),
        document.get('expected_evidence') or '',
        document.get('primary_purpose') or '',
    )


def _stored_catalog(db: Session, version_id: int) -> tuple[dict, dict]:
    document_keys = {
        item.id: _norm(item.document_type)
        for item in db.scalars(select(DocumentTypeRule).where(
            DocumentTypeRule.checklist_version_id == version_id
        )).all()
    }
    rules = {
        item.rule_id: {
            'practice_area': item.practice_area_code,
            'level': item.level,
            'audit_check': item.audit_check,
            'gap_text': item.gap_text,
            'document_key': document_keys.get(item.document_type_rule_id),
            'condition': item.applicability_condition,
            'detection': item.detection,
            'recommendation': item.recommendation,
            'is_mandatory': item.is_mandatory,
        }
        for item in db.scalars(select(CmmiRule).where(CmmiRule.checklist_version_id == version_id)).all()
    }
    documents = {
        item.document_type: {
            'practice_areas': item.practice_areas or [],
            'keywords': item.keywords or [],
            'aliases': item.aliases or [],
            'expected_evidence': item.expected_evidence or '',
            'primary_purpose': item.primary_purpose or '',
        }
        for item in db.scalars(select(DocumentTypeRule).where(DocumentTypeRule.checklist_version_id == version_id)).all()
    }
    return rules, documents


def _change_summary(db: Session, active: ChecklistVersion | None, rules: list[dict], documents: list[dict]) -> dict:
    """Describe substantive rulebook changes without exposing workbook bytes."""
    incoming_rules = {item['rule_id']: item for item in rules}
    incoming_documents = {item['document_type']: item for item in documents}
    if not active:
        return {
            'rules_added': len(incoming_rules), 'rules_removed': 0, 'rules_changed': 0,
            'documents_added': len(incoming_documents), 'documents_removed': 0, 'documents_changed': 0,
            'changed_rule_ids': sorted(incoming_rules)[:25], 'changed_document_types': sorted(incoming_documents)[:25],
        }
    stored_rules, stored_documents = _stored_catalog(db, active.id)
    changed_rules = sorted(rule_id for rule_id in incoming_rules.keys() & stored_rules.keys()
                           if _rule_signature(incoming_rules[rule_id]) != _rule_signature(stored_rules[rule_id]))
    changed_documents = sorted(name for name in incoming_documents.keys() & stored_documents.keys()
                               if _document_signature(incoming_documents[name]) != _document_signature(stored_documents[name]))
    return {
        'rules_added': len(incoming_rules.keys() - stored_rules.keys()),
        'rules_removed': len(stored_rules.keys() - incoming_rules.keys()),
        'rules_changed': len(changed_rules),
        'documents_added': len(incoming_documents.keys() - stored_documents.keys()),
        'documents_removed': len(stored_documents.keys() - incoming_documents.keys()),
        'documents_changed': len(changed_documents),
        'changed_rule_ids': changed_rules[:25],
        'changed_document_types': changed_documents[:25],
    }


def _has_substantive_change(summary: dict) -> bool:
    return any(summary[key] for key in (
        'rules_added', 'rules_removed', 'rules_changed',
        'documents_added', 'documents_removed', 'documents_changed',
    ))


def activate_rule_catalog(db: Session, checklist_version_id: int) -> ChecklistVersion:
    """Promote a validated immutable ruleset only through an explicit action."""
    # Lock the complete, small version registry so two concurrent activation
    # requests cannot both leave their independently selected rows active.
    versions = db.scalars(select(ChecklistVersion).order_by(
        ChecklistVersion.id
    ).with_for_update()).all()
    version = next((item for item in versions if item.id == checklist_version_id), None)
    if not version:
        raise ValueError('Ruleset version not found.')
    if version.status not in {'VALIDATED', 'ACTIVE', 'RETIRED'}:
        raise ValueError('Only a validated or previously governed ruleset can be activated.')
    if version.is_active:
        return version
    db.execute(update(ChecklistVersion).where(ChecklistVersion.is_active.is_(True)).values(
        is_active=False, status='RETIRED'
    ))
    version.is_active = True
    version.status = 'ACTIVE'
    return version


def import_rule_catalog(db: Session, filename: str, content: bytes, activate: bool = False) -> dict:
    """Validate and version a CMMI or final SAP evidence catalogue.

    A SAP workbook is identified by its two required sheets and creates a
    version containing the same fixed 14 evidence types plus its 210
    document-linked field controls.  CMMI workbook behaviour is unchanged.
    """
    if not filename.lower().endswith(('.xlsx', '.xlsm')):
        raise ValueError('Master rules must be supplied as an XLSX workbook.')
    digest = sha256(content).hexdigest()
    existing = db.scalar(select(ChecklistVersion).where(ChecklistVersion.checksum == digest))
    if existing:
        if activate and not existing.is_active:
            activate_rule_catalog(db, existing.id)
        return {'checklist_version_id': existing.id, 'version': existing.version, 'rules': db.query(CmmiRule).filter_by(checklist_version_id=existing.id).count(),
                'document_types': db.query(DocumentTypeRule).filter_by(checklist_version_id=existing.id).count(),
                'already_imported': True, 'content_unchanged': True, 'change_summary': {}}
    workbook = _load_safe_workbook(content, 'The master workbook is corrupt, unsafe, or unreadable.')
    try:
        is_sap_checklist = 'Master Checklist' in workbook.sheetnames
        if is_sap_checklist:
            known_practice_areas = set(db.scalars(select(PracticeArea.code)).all())
            documents, rules = _parse_sap_audit_checklist(workbook, known_practice_areas)
            framework = 'SAP Evidence Audit'
            version_prefix = 'SAP-AUDIT'
        else:
            rules = _parse_rules(workbook)
            # Keep document matching/AFR scope fixed to the approved 14
            # entries for the original CMMI workbook profile.
            documents = _approved_documents()
            framework = 'CMMI v3.0'
            version_prefix = 'CMMI-V3'
    finally:
        workbook.close()
    active = db.scalar(select(ChecklistVersion).where(ChecklistVersion.is_active.is_(True)).order_by(ChecklistVersion.id.desc()))
    summary = _change_summary(db, active, rules, documents)
    if active and not _has_substantive_change(summary):
        return {
            'checklist_version_id': active.id, 'version': active.version,
            'rules': len(rules), 'document_types': len(documents), 'already_imported': True,
            'content_unchanged': True, 'change_summary': summary,
        }
    version_name = f'{version_prefix}-{date.today().isoformat()}-{digest[:8]}'
    version = ChecklistVersion(version=version_name, source=filename[:255], checksum=digest, framework=framework,
                               status='VALIDATED', effective_from=date.today(), is_active=False)
    db.add(version); db.flush()
    db.add_all([DocumentTypeRule(checklist_version_id=version.id, **item) for item in documents])
    db.flush()
    document_ids = {
        _norm(item.document_type): item.id
        for item in db.scalars(select(DocumentTypeRule).where(
            DocumentTypeRule.checklist_version_id == version.id
        )).all()
    }
    unmapped_rule_ids = [
        item['rule_id'] for item in rules
        if item.get('document_key') and item['document_key'] not in document_ids
    ]
    if unmapped_rule_ids:
        raise RuntimeError(
            'A validated rule could not be linked to its evidence type: '
            + ', '.join(unmapped_rule_ids[:5])
        )
    db.add_all([CmmiRule(
        checklist_version_id=version.id, rule_id=item['rule_id'], practice_area_code=item['practice_area'],
        level=item['level'], audit_check=item['audit_check'], gap_text=item['gap_text'],
        document_type_rule_id=document_ids.get(item.get('document_key')),
        applicability_condition=item.get('condition'), detection=item.get('detection'),
        recommendation=item.get('recommendation'), is_mandatory=item.get('is_mandatory', True),
    ) for item in rules])
    if activate:
        activate_rule_catalog(db, version.id)
    return {'checklist_version_id': version.id, 'version': version.version, 'rules': len(rules),
            'document_types': len(documents), 'already_imported': False,
            'content_unchanged': False, 'change_summary': summary}


def import_document_catalogue_additions(db: Session, filename: str, content: bytes) -> dict:
    """Merge keys into an approved active document type without widening scope.

    A catalogue worksheet may improve the configured keys/evidence wording
    for one of the approved 14 types. It cannot create another type or change
    a type's practice-area mapping, as either action would widen CMMI and AFR
    scope without governance.
    """
    if not filename.lower().endswith('.xlsx'):
        raise ValueError('Document catalogue additions must be supplied as an XLSX workbook.')
    workbook = _load_safe_workbook(content, 'The document catalogue workbook is corrupt, unsafe, or unreadable.')
    known_practice_areas = set(db.scalars(select(PracticeArea.code)).all())
    try:
        additions = _parse_document_additions(workbook, known_practice_areas)
    finally:
        workbook.close()
    # Serialize additions for the active catalogue. This prevents two uploads
    # from independently inserting the same normalized document type.
    active_versions = db.scalars(select(ChecklistVersion).where(
        ChecklistVersion.is_active.is_(True)
    ).order_by(ChecklistVersion.id.desc()).with_for_update()).all()
    if not active_versions:
        raise ValueError('No active ruleset is available for document catalogue additions.')
    if len(active_versions) != 1:
        raise ValueError('Ruleset state is invalid: exactly one active version is required.')
    active = active_versions[0]
    stored_documents = db.scalars(select(DocumentTypeRule).where(
        DocumentTypeRule.checklist_version_id == active.id
    )).all()
    existing: dict[str, DocumentTypeRule] = {}
    for item in stored_documents:
        key = _norm(item.document_type)
        if key in existing:
            raise ValueError(
                'The active ruleset contains ambiguous document names: '
                f'"{existing[key].document_type}" and "{item.document_type}". '
                'Resolve the existing duplicate before importing additions.'
            )
        existing[key] = item
    approved = {_norm(item['document_type']): item for item in _approved_documents()}
    if set(existing) != set(approved):
        raise ValueError(
            'The active ruleset document catalogue is not the governed 14-entry catalogue. '
            'Create or activate a validated ruleset before updating document keys.'
        )

    # Validate every row before mutating ORM objects. This makes a mixed
    # worksheet all-or-nothing even when the caller handles the exception.
    for addition in additions:
        key = _norm(addition['document_type'])
        document = existing.get(key)
        if not document:
            raise ValueError(
                f'"{addition["document_type"]}" is not one of the approved 14 document types. '
                'New document types cannot be added through this upload.'
            )
        unexpected_practice_areas = set(addition['practice_areas']) - set(document.practice_areas or [])
        if unexpected_practice_areas:
            raise ValueError(
                f'"{document.document_type}" has a fixed practice-area mapping. '
                f'Unsupported mapping(s): {", ".join(sorted(unexpected_practice_areas))}.'
            )
    inserted = updated = unchanged = 0
    for addition in additions:
        document = existing.get(_norm(addition['document_type']))
        merged_keywords = _merge_text_values(document.keywords or [], addition['keywords'])
        merged_aliases = _merge_text_values(document.aliases or [], addition['aliases'])
        merged_expected = '; '.join(_merge_text_values(
            _split_expected_evidence(document.expected_evidence), addition['expected_evidence']
        ))
        changed = (
            merged_keywords != (document.keywords or [])
            or merged_aliases != (document.aliases or [])
            or merged_expected != (document.expected_evidence or '')
            or not document.include_in_afr
        )
        if not changed:
            unchanged += 1
            continue
        document.keywords = merged_keywords
        document.aliases = merged_aliases
        document.expected_evidence = merged_expected
        # The approved 14 entries are the explicitly approved AFR catalogue.
        document.include_in_afr = True
        updated += 1
    return {
        'checklist_version_id': active.id,
        'version': active.version,
        'source_checksum': sha256(content).hexdigest(),
        'rows_received': len(additions),
        'document_types_inserted': inserted,
        'document_types_updated': updated,
        'document_types_unchanged': unchanged,
    }
