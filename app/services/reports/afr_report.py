from datetime import date, datetime, time, timezone
from pathlib import Path
from uuid import uuid4
from sqlalchemy import select
from sqlalchemy.orm import Session
from app.core.config import get_settings
from app.db.models import AuditProject, AuditSession, Comment, EvidenceFile, EvidenceSource, Finding, GeneratedReport, RemediationAction
from app.services.finding_scope import afr_findings_query, evidence_findings_query
from app.services.reports.finding_export import (
    DETAILED_CONTROL_CONTEXT_HEADERS, add_empty_findings_message, audit_scope_status,
    classification_display_values, evidence_link_status, finding_evidence_paths,
    display_finding_status, display_rule_id, format_available_keys, format_key_values, safe_export_value, style_classification_sheet,
    CLASSIFICATION_HEADERS, _append_document_coverage_sheet,
)

def create_afr(db: Session, audit_session_id: int, user_id: int, fmt: str,
               practice_area_codes: list[str] | None = None,
               severities: list[str] | None = None,
               statuses: list[str] | None = None,
               from_date: date | None = None, to_date: date | None = None) -> GeneratedReport:
    """Create an AFR from approved-catalogue findings and optional filters."""
    audit = db.get(AuditSession, audit_session_id)
    if not audit: raise ValueError('Audit session not found')
    audit_context = {
        'project_name': db.scalar(select(AuditProject.name).where(AuditProject.id == audit.project_id)) or '',
        'name': audit.audit_name or f'Audit session #{audit_session_id}',
        'date': audit.audit_date.isoformat() if audit.audit_date else '',
        'auditors': audit.auditors or '',
        'auditees': audit.auditees or '',
    }
    # An AFR is a current report by default.  A caller may explicitly request
    # historical statuses (for example, superseded findings) through `statuses`.
    query = afr_findings_query(audit_session_id, open_only=not statuses)
    if practice_area_codes: query = query.where(Finding.practice_area_code.in_(practice_area_codes))
    if severities: query = query.where(Finding.severity.in_(severities))
    if statuses: query = query.where(Finding.status.in_(statuses))
    if from_date: query = query.where(Finding.created_at >= datetime.combine(from_date, time.min, tzinfo=timezone.utc))
    if to_date: query = query.where(Finding.created_at <= datetime.combine(to_date, time.max, tzinfo=timezone.utc))
    findings = db.scalars(query.order_by(Finding.severity, Finding.practice_area_code, Finding.id)).all()
    finding_ids = [finding.id for finding in findings]
    comments = {finding_id: [] for finding_id in finding_ids}
    actions = {finding_id: [] for finding_id in finding_ids}
    evidence_by_finding = finding_evidence_paths(
        db, audit_session_id, finding_ids, afr_eligible_only=True
    )
    scope_status, unreadable_paths = audit_scope_status(db, audit_session_id)
    if finding_ids:
        for finding_id, body in db.execute(
            select(Comment.finding_id, Comment.body)
            .where(Comment.finding_id.in_(finding_ids))
            .order_by(Comment.finding_id, Comment.created_at)
        ):
            comments[finding_id].append(body)
        for finding_id, status, action in db.execute(
            select(RemediationAction.finding_id, RemediationAction.status, RemediationAction.action)
            .where(RemediationAction.finding_id.in_(finding_ids))
            .order_by(RemediationAction.finding_id, RemediationAction.created_at)
        ):
            actions[finding_id].append(f'{status}: {action}')
    created = datetime.now(timezone.utc)
    output = get_settings().output_dir / f'AFR_session-{audit_session_id}_{created:%Y%m%dT%H%M%SZ}_{uuid4().hex[:8]}.{fmt}'
    if fmt == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        wb=Workbook(); ws=wb.active; ws.title='Detailed Findings'
        ws.append(['Finding ID','Rule ID','Practice Area','Severity','Finding','Required Keys','Available Keys','Missing Required Keys','Evidence Link Status','Evidence File(s)','Description','Recommendation','Status','Comments','Remediation actions'])
        for f in findings:
            paths = evidence_by_finding[f.id]
            ws.append([safe_export_value(value) for value in (
                f.id, display_rule_id(f.rule_id), f.practice_area_code, f.severity, f.title,
                format_key_values(f.required_keys), format_available_keys(
                    f.available_keys, has_linked_evidence=bool(paths)
                ), format_key_values(f.missing_required_keys),
                evidence_link_status(paths, unreadable_paths), '\n'.join(paths), f.description, f.recommendation,
                display_finding_status(f), '\n'.join(comments[f.id]), '\n'.join(actions[f.id]),
            )])
        ws.freeze_panes = 'A2'
        header_fill = PatternFill('solid', fgColor='1F4E78')
        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.row_dimensions[1].height = 32
        ws.auto_filter.ref = f'A1:O{max(1, ws.max_row)}'
        widths = [12, 14, 15, 12, 34, 30, 30, 32, 30, 42, 56, 56, 14, 36, 36]
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(1, index).column_letter].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
            row[3].alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)
            row[12].alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)
        if not findings:
            add_empty_findings_message(ws, 15)
        # Preserve the AFR summary on the first worksheet and place every
        # file-backed field-control result on a second worksheet in the same
        # workbook. This makes the summary concise without hiding scan detail.
        control_query = evidence_findings_query(audit_session_id, open_only=not statuses).where(
            Finding.finding_kind != 'document_catalogue'
        )
        if practice_area_codes: control_query = control_query.where(Finding.practice_area_code.in_(practice_area_codes))
        if severities: control_query = control_query.where(Finding.severity.in_(severities))
        if statuses: control_query = control_query.where(Finding.status.in_(statuses))
        if from_date: control_query = control_query.where(Finding.created_at >= datetime.combine(from_date, time.min, tzinfo=timezone.utc))
        if to_date: control_query = control_query.where(Finding.created_at <= datetime.combine(to_date, time.max, tzinfo=timezone.utc))
        control_findings = db.scalars(control_query.order_by(
            Finding.severity, Finding.practice_area_code, Finding.id
        )).all()
        control_ids = [finding.id for finding in control_findings]
        control_evidence = finding_evidence_paths(db, audit_session_id, control_ids, afr_eligible_only=True)
        control_comments = {finding_id: [] for finding_id in control_ids}
        control_actions = {finding_id: [] for finding_id in control_ids}
        if control_ids:
            for finding_id, body in db.execute(select(Comment.finding_id, Comment.body).where(
                Comment.finding_id.in_(control_ids)
            ).order_by(Comment.finding_id, Comment.created_at)):
                control_comments[finding_id].append(body)
            for finding_id, status, action in db.execute(select(
                RemediationAction.finding_id, RemediationAction.status, RemediationAction.action
            ).where(RemediationAction.finding_id.in_(control_ids)).order_by(
                RemediationAction.finding_id, RemediationAction.created_at)):
                control_actions[finding_id].append(f'{status}: {action}')
        detail = wb.create_sheet('Detailed Control Findings')
        detail.append([
            *DETAILED_CONTROL_CONTEXT_HEADERS,
            'Finding ID', 'Rule ID', 'Practice Area', 'Severity', 'Finding',
            'Required Keys', 'Available Keys', 'Missing Required Keys', 'Evidence Link Status',
            'Evidence File(s)', 'Description', 'Recommendation', 'Status', 'Comments', 'Remediation actions',
        ])
        for f in control_findings:
            paths = control_evidence[f.id]
            detail.append([safe_export_value(value) for value in (
                audit_context['project_name'], audit_context['date'], audit_context['auditors'], audit_context['auditees'],
                f.id, display_rule_id(f.rule_id), f.practice_area_code, f.severity, f.title,
                format_key_values(f.required_keys), format_available_keys(
                    f.available_keys, has_linked_evidence=bool(paths)
                ), format_key_values(f.missing_required_keys),
                evidence_link_status(paths, unreadable_paths), '\n'.join(paths), f.description, f.recommendation,
                display_finding_status(f), '\n'.join(control_comments[f.id]), '\n'.join(control_actions[f.id]),
            )])
        detail.freeze_panes = 'A2'
        detail.auto_filter.ref = f'A1:{detail.cell(1, detail.max_column).column_letter}{max(1, detail.max_row)}'
        for cell in detail[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        detail_widths = [28, 16, 28, 28] + widths
        for index, width in enumerate(detail_widths, start=1):
            detail.column_dimensions[detail.cell(1, index).column_letter].width = width
        for row in detail.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        classification = wb.create_sheet('Evidence Classification')
        classification.append(CLASSIFICATION_HEADERS)
        evidence_files = db.scalars(select(EvidenceFile).join(EvidenceSource).where(
            EvidenceSource.audit_session_id == audit_session_id,
        ).order_by(EvidenceFile.relative_path, EvidenceFile.id)).all()
        for evidence in evidence_files:
            result = evidence.classification_json or {}
            classification.append([safe_export_value(value) for value in classification_display_values(
                evidence.relative_path, result
            )])
        classification.freeze_panes = 'A2'
        style_classification_sheet(classification)
        _append_document_coverage_sheet(wb, db, audit_session_id)
        meta=wb.create_sheet('Report Metadata'); meta.append(['Audit Session',audit_session_id]); meta.append(['Audit name',audit_context['name']]); meta.append(['Audit date',audit_context['date']]); meta.append(['Auditors',audit_context['auditors']]); meta.append(['Auditees',audit_context['auditees']]); meta.append(['Generated At',created.isoformat()]); meta.append(['Generated By',user_id]); meta.append(['Practice areas', ', '.join(practice_area_codes or []) or 'All']); meta.append(['Severities', ', '.join(severities or []) or 'All']); meta.append(['Statuses', ', '.join(statuses or []) or 'All']); meta.append(['From date', from_date.isoformat() if from_date else '']); meta.append(['To date', to_date.isoformat() if to_date else '']); meta.append(['Evidence assessment scope',scope_status]); meta.append(['Unreadable evidence file(s)','\n'.join(unreadable_paths) or 'None']); meta.append(['Report scope','Detailed Findings contains confirmed file-backed findings and project-scope coverage gaps. Detailed Control Findings contains the file-backed auditor-review queue. Evidence Classification records the document mapping used for this scan.'])
        meta.column_dimensions['A'].width = 32
        meta.column_dimensions['B'].width = 110
        for label, value in meta.iter_rows():
            label.font = Font(bold=True, color='FFFFFF')
            label.fill = header_fill
            label.alignment = Alignment(vertical='top', wrap_text=True)
            value.alignment = Alignment(vertical='top', wrap_text=True)
        wb.save(output)
    elif fmt == 'docx':
        from docx import Document
        doc=Document(); doc.add_heading('CMMI Audit Findings Report (AFR)',0); doc.add_paragraph(f"Audit: {audit_context['name']} | Session: {audit_session_id} | Date: {audit_context['date'] or 'Not recorded'} | Auditors: {audit_context['auditors'] or 'Not recorded'} | Auditees: {audit_context['auditees'] or 'Not recorded'} | Generated: {created.isoformat()} | Creator: {user_id}"); doc.add_paragraph('Scope: confirmed findings linked to approved AFR evidence and project-scope coverage gaps are included. Auditor-review items remain in the detailed control output.')
        table=doc.add_table(rows=1, cols=4); table.style='Table Grid'
        for cell,value in zip(table.rows[0].cells,['Rule','PA','Severity','Finding']): cell.text=value
        for f in findings:
            cells=table.add_row().cells
            for cell,value in zip(cells,[display_rule_id(f.rule_id),f.practice_area_code,f.severity,f.title]): cell.text=str(value)
            paths = evidence_by_finding[f.id]
            doc.add_paragraph(f'Evidence trace: {evidence_link_status(paths, unreadable_paths)}' + (f" — {' | '.join(paths)}" if paths else ''))
            doc.add_paragraph(f'Recommendation: {f.recommendation}')
            if comments[f.id]: doc.add_paragraph(f'Comments: {" | ".join(comments[f.id])}')
            if actions[f.id]: doc.add_paragraph(f'Remediation: {" | ".join(actions[f.id])}')
        doc.save(output)
    else:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        pdf=canvas.Canvas(str(output), pagesize=A4); y=800; pdf.setFont('Helvetica-Bold',16); pdf.drawString(48,y,'CMMI Audit Findings Report (AFR)'); y-=26; pdf.setFont('Helvetica',8)
        pdf.drawString(48,y,f"Audit: {audit_context['name']} | session {audit_session_id} | date: {audit_context['date'] or 'Not recorded'}"); y-=12
        pdf.drawString(48,y,f"Auditors: {audit_context['auditors'] or 'Not recorded'} | Auditees: {audit_context['auditees'] or 'Not recorded'}"); y-=12
        pdf.drawString(48,y,f'Generated {created.isoformat()} by {user_id}'); y-=12
        pdf.drawString(48,y,'Scope: confirmed approved-evidence findings and project coverage gaps.'); y-=12
        for f in findings:
            line=f'[{f.severity.upper()}] {f.practice_area_code} {display_rule_id(f.rule_id)}: {f.title}'
            for segment in [line[i:i+105] for i in range(0,len(line),105)]:
                if y<50: pdf.showPage(); y=800
                pdf.drawString(48,y,segment); y-=12
            paths = evidence_by_finding[f.id]
            trace = f'Evidence: {evidence_link_status(paths, unreadable_paths)}' + (f" — {' | '.join(paths)}" if paths else '')
            for segment in [trace[i:i+105] for i in range(0,len(trace),105)]:
                if y<50: pdf.showPage(); y=800
                pdf.drawString(48,y,segment); y-=12
        pdf.save()
    report=GeneratedReport(audit_session_id=audit_session_id,report_type=f'afr_{fmt}',storage_path=str(output),created_by_id=user_id); db.add(report); db.commit(); db.refresh(report); return report
