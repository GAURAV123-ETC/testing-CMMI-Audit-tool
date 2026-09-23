"""Exports the current audit session's persisted findings for Evidence Scan."""
import csv
from datetime import datetime, timezone
from io import StringIO
from pathlib import Path
from uuid import uuid4

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import get_settings
from app.db.models import (AuditProject, AuditSession, DocumentTypeRule, EvidenceFile,
                           EvidenceSource, Finding, FindingEvidence, GeneratedReport)
from app.services.finding_scope import afr_findings_query, evidence_findings_query


HEADERS = ('Finding ID', 'Rule ID', 'Practice Area', 'Severity', 'Status', 'Finding',
           'Description', 'Recommendation', 'Required Keys', 'Available Keys', 'Missing Required Keys',
           'Evidence Link Status', 'Evidence File(s)', 'Created At')
DETAILED_CONTROL_CONTEXT_HEADERS = ('Project Name', 'Audit Date', 'Auditor(s)', 'Auditee(s)')
DETAILED_CONTROL_HEADERS = (*DETAILED_CONTROL_CONTEXT_HEADERS, *HEADERS)
EXCEL_CELL_LIMIT = 32_767
NO_AVAILABLE_KEYS_IN_DOCUMENT = 'No reliably matched available key was found in this document.'
NO_AVAILABLE_KEYS_IN_SCOPE = 'No matching document or available key was found in the audit scope.'
NO_CONFIRMED_FINDINGS = (
    'No confirmed open findings. Controls requiring auditor review are listed in '
    'Detailed Control Findings.'
)
DOCUMENT_COVERAGE_RULE_LABEL = 'Document Coverage'
CLASSIFICATION_HEADERS = (
    'Evidence File', 'Detected Document Type', 'Practice Area(s)', 'Confidence',
    'Matched Keys', 'Missing Keys', 'Classification Reason', 'Document Role',
)
DOCUMENT_COVERAGE_HEADERS = (
    'Checklist Document Type', 'Practice Area(s)', 'Expected Evidence',
    'Upload / Match Status', 'Matched Uploaded File(s)', 'Detected As',
    'Classification Confidence', 'Assessment Note',
)
EXCLUDED_DOCUMENT_ROLES = {
    'TEMPLATE', 'BLANK_TEMPLATE', 'PROCESS_REFERENCE', 'DUPLICATE', 'SUPERSEDED_VERSION',
}


def safe_export_value(value: object) -> str | int:
    """Prevent spreadsheet formula execution and oversized/corrupt cells."""
    if isinstance(value, int):
        return value
    text = str(value or '').replace('\x00', '')
    # Excel/CSV interprets these leading characters as a formula when opened.
    if text.startswith(('=', '+', '-', '@')):
        text = "'" + text
    if len(text) > EXCEL_CELL_LIMIT:
        text = text[:EXCEL_CELL_LIMIT - 15] + ' [truncated]'
    return text


def format_key_values(values: object) -> str:
    """Render JSON key values safely, including defensively handled legacy rows."""
    if values is None:
        candidates = []
    elif isinstance(values, (str, bytes)):
        candidates = [values]
    elif isinstance(values, (list, tuple, set)):
        candidates = values
    else:
        candidates = [values]
    return '\n'.join(str(value).strip() for value in candidates if value is not None and str(value).strip())


def display_rule_id(rule_id: object) -> str:
    """Keep internal coverage identifiers out of end-user report columns."""
    value = str(rule_id or '').strip()
    return DOCUMENT_COVERAGE_RULE_LABEL if value.startswith('DOC-COVERAGE-') else value


def format_available_keys(values: object, *, has_linked_evidence: bool) -> str:
    """Render an explicit absence reason instead of an ambiguous blank cell."""
    rendered = format_key_values(values)
    if rendered:
        return rendered
    return NO_AVAILABLE_KEYS_IN_DOCUMENT if has_linked_evidence else NO_AVAILABLE_KEYS_IN_SCOPE


def add_empty_findings_message(sheet, last_column: int) -> None:
    """Make an intentionally empty AFR summary visibly distinguishable from a failed export."""
    from openpyxl.styles import Alignment, Font

    sheet.cell(row=2, column=1, value=NO_CONFIRMED_FINDINGS)
    if last_column > 1:
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    sheet.cell(row=2, column=1).alignment = Alignment(vertical='top', wrap_text=True)
    sheet.cell(row=2, column=1).font = Font(italic=True, color='666666')


def classification_display_values(evidence_path: str, result: dict) -> list[str]:
    """Make classification states explicit instead of leaving report cells blank."""
    role = str(result.get('document_role') or '').strip()
    excluded = role in EXCLUDED_DOCUMENT_ROLES
    practice_areas = result.get('practice_areas') or []
    matched_keys = result.get('matched_keywords') or []
    missing_keys = result.get('missed_keywords') or []
    return [
        evidence_path,
        str(result.get('detected_type') or 'Not scanned'),
        '\n'.join(practice_areas) or (
            'Not applicable - reference document' if excluded else 'Not classified'
        ),
        str(result.get('confidence') or ('Not applicable' if excluded else 'Not assessed')),
        '\n'.join(matched_keys) or (
            'Not applicable - reference document' if excluded else 'No configured document keys matched'
        ),
        '\n'.join(missing_keys) or (
            'Not applicable - reference document' if excluded else 'No missing classification keys identified'
        ),
        str(result.get('classification_reason') or 'No classification explanation recorded.'),
        role or 'Not recorded',
    ]


def style_classification_sheet(sheet) -> None:
    """Keep key lists readable in Excel instead of visually concatenating them."""
    from openpyxl.styles import Alignment, Font, PatternFill

    header_fill = PatternFill('solid', fgColor='1F4E78')
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    widths = [52, 40, 18, 14, 38, 38, 78, 30]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    sheet.row_dimensions[1].height = 30
    for row in sheet.iter_rows(min_row=2):
        line_count = max((str(cell.value or '').count('\n') + 1 for cell in row), default=1)
        sheet.row_dimensions[row[0].row].height = min(90, max(24, 15 * line_count))
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)


def finding_evidence_paths(db: Session, audit_session_id: int, finding_ids: list[int],
                           *, afr_eligible_only: bool = False) -> dict[int, list[str]]:
    """Load session-scoped direct evidence links, preserving stable file paths."""
    evidence_by_finding: dict[int, list[str]] = {finding_id: [] for finding_id in finding_ids}
    if not finding_ids:
        return evidence_by_finding
    query = select(FindingEvidence.finding_id, EvidenceFile.relative_path).join(
        Finding, Finding.id == FindingEvidence.finding_id
    ).join(
        EvidenceFile, EvidenceFile.id == FindingEvidence.evidence_file_id
    ).join(
        EvidenceSource, EvidenceSource.id == EvidenceFile.source_id
    ).where(FindingEvidence.finding_id.in_(finding_ids)).order_by(
        FindingEvidence.finding_id, EvidenceFile.relative_path
    ).where(
        Finding.audit_session_id == audit_session_id,
        EvidenceSource.audit_session_id == audit_session_id,
    )
    if afr_eligible_only:
        from sqlalchemy import true
        query = query.where(EvidenceFile.classification_json['afr_eligible'].as_boolean() == true())
    links = db.execute(query).all()
    for finding_id, path in links:
        evidence_by_finding[finding_id].append(path)
    return evidence_by_finding


def audit_scope_status(db: Session, audit_session_id: int) -> tuple[str, list[str]]:
    """Identify unreadable evidence that makes absence-based gaps provisional."""
    unreadable_paths = list(db.scalars(select(EvidenceFile.relative_path).join(EvidenceSource).where(
        EvidenceSource.audit_session_id == audit_session_id,
        EvidenceFile.processing_status == 'error',
    ).order_by(EvidenceFile.relative_path)).all())
    if unreadable_paths:
        return (
            f'Incomplete: {len(unreadable_paths)} uploaded file(s) could not be assessed',
            unreadable_paths,
        )
    return 'Complete: all uploaded files were assessed', []


def evidence_link_status(paths: list[str], unreadable_paths: list[str] | None = None) -> str:
    if paths:
        base = 'Uploaded evidence assessed for this finding'
    else:
        base = 'No matching uploaded evidence in this audit-session scope'
    if unreadable_paths:
        return f'{base}; assessment scope incomplete ({len(unreadable_paths)} unreadable file(s))'
    return base


def _rows(db: Session, audit_session_id: int, findings: list[Finding] | None = None) -> list[list[str | int]]:
    if findings is None:
        findings = db.scalars(afr_findings_query(audit_session_id).order_by(
            Finding.severity, Finding.practice_area_code, Finding.id
        )).all()
    evidence_by_finding = finding_evidence_paths(
        db, audit_session_id, [finding.id for finding in findings], afr_eligible_only=True
    )
    _, unreadable_paths = audit_scope_status(db, audit_session_id)
    return [[safe_export_value(value) for value in (
        finding.id, display_rule_id(finding.rule_id), finding.practice_area_code, finding.severity,
        finding.status, finding.title, finding.description, finding.recommendation,
        format_key_values(finding.required_keys), format_available_keys(
            finding.available_keys, has_linked_evidence=bool(evidence_by_finding[finding.id])
        ),
        format_key_values(finding.missing_required_keys),
        evidence_link_status(evidence_by_finding[finding.id], unreadable_paths), '\n'.join(evidence_by_finding[finding.id]),
        finding.created_at.isoformat(),
    )] for finding in findings]


def _detailed_control_context(db: Session, audit_session_id: int) -> list[str]:
    """Return audit metadata repeated beside each detailed control finding."""
    audit = db.get(AuditSession, audit_session_id)
    if not audit:
        raise ValueError('Audit session not found')
    project_name = db.scalar(select(AuditProject.name).where(AuditProject.id == audit.project_id)) or ''
    return [
        safe_export_value(project_name),
        safe_export_value(audit.audit_date.isoformat() if audit.audit_date else ''),
        safe_export_value(audit.auditors),
        safe_export_value(audit.auditees),
    ]


def _append_findings_sheet(workbook, title: str, rows: list[list[str | int]],
                           headers: tuple[str, ...] = HEADERS) -> None:
    sheet = workbook.create_sheet(title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = 'A2'
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(60, max(14, max(len(str(cell.value or '')) for cell in column) + 2))


def _append_classification_sheet(workbook, db: Session, audit_session_id: int) -> None:
    sheet = workbook.create_sheet('Evidence Classification')
    sheet.append(CLASSIFICATION_HEADERS)
    files = db.scalars(select(EvidenceFile).join(EvidenceSource).where(
        EvidenceSource.audit_session_id == audit_session_id,
    ).order_by(EvidenceFile.relative_path, EvidenceFile.id)).all()
    for evidence in files:
        result = evidence.classification_json or {}
        sheet.append([safe_export_value(value) for value in classification_display_values(
            evidence.relative_path, result
        )])
    sheet.freeze_panes = 'A2'
    style_classification_sheet(sheet)
    coverage_widths = [40, 18, 62, 30, 56, 36, 24, 64]
    for index, width in enumerate(coverage_widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width


def _append_document_coverage_sheet(workbook, db: Session, audit_session_id: int) -> None:
    """Add a catalogue-level view of exactly what the scan did and did not match.

    The Evidence Classification sheet is file-centric.  This complementary
    sheet is checklist-centric, so an auditor can see every governed document
    type even when no corresponding upload exists.  It relies only on the
    persisted scan classification ID; it never re-classifies from a filename
    while exporting the AFR.
    """
    audit = db.get(AuditSession, audit_session_id)
    if not audit:
        raise ValueError('Audit session not found')
    sheet = workbook.create_sheet('Document Coverage')
    sheet.append(DOCUMENT_COVERAGE_HEADERS)
    document_types = db.scalars(select(DocumentTypeRule).where(
        DocumentTypeRule.checklist_version_id == audit.checklist_version_id,
        DocumentTypeRule.include_in_afr.is_(True),
    ).order_by(DocumentTypeRule.document_type, DocumentTypeRule.id)).all()
    evidence_files = db.scalars(select(EvidenceFile).join(EvidenceSource).where(
        EvidenceSource.audit_session_id == audit_session_id,
    ).order_by(EvidenceFile.relative_path, EvidenceFile.id)).all()

    matches_by_document_id: dict[int, list[EvidenceFile]] = {}
    for evidence in evidence_files:
        classification = evidence.classification_json or {}
        document_type_id = classification.get('document_type_rule_id')
        if isinstance(document_type_id, int):
            matches_by_document_id.setdefault(document_type_id, []).append(evidence)

    for document_type in document_types:
        matches = matches_by_document_id.get(document_type.id, [])
        implementation_matches = [evidence for evidence in matches if (
            (evidence.classification_json or {}).get('document_role') == 'PROJECT_IMPLEMENTATION_EVIDENCE'
            and evidence.processing_status == 'processed'
        )]
        if implementation_matches:
            status = 'Matched'
            files = '\n'.join(evidence.relative_path for evidence in implementation_matches)
            detected_as = '\n'.join(dict.fromkeys(
                str((evidence.classification_json or {}).get('detected_type') or 'Not recorded')
                for evidence in implementation_matches
            ))
            confidence = '\n'.join(dict.fromkeys(
                str((evidence.classification_json or {}).get('confidence') or 'Not assessed')
                for evidence in implementation_matches
            ))
            note = f'{len(implementation_matches)} uploaded file(s) matched this checklist document type.'
        elif matches:
            status = 'Uploaded but not assessable'
            files = '\n'.join(evidence.relative_path for evidence in matches)
            detected_as = '\n'.join(dict.fromkeys(
                str((evidence.classification_json or {}).get('detected_type') or 'Not recorded')
                for evidence in matches
            ))
            confidence = '\n'.join(dict.fromkeys(
                str((evidence.classification_json or {}).get('confidence') or 'Not assessed')
                for evidence in matches
            ))
            note = 'A file mapped to this type, but it was not usable as project implementation evidence. Review Evidence Classification.'
        else:
            status = 'Not matched in uploaded evidence'
            files = 'No uploaded file matched this checklist document type.'
            detected_as = 'Not applicable'
            confidence = 'Not applicable'
            note = 'Upload the governed document if it is in scope for this audit.'
        sheet.append([safe_export_value(value) for value in (
            document_type.document_type,
            '\n'.join(document_type.practice_areas or []) or 'Not recorded',
            document_type.expected_evidence or 'No expected evidence description recorded.',
            status, files, detected_as, confidence, note,
        )])
    sheet.freeze_panes = 'A2'
    style_classification_sheet(sheet)


def create_current_findings_export(db: Session, audit_session_id: int, user_id: int, fmt: str) -> GeneratedReport:
    """Create a CSV or XLSX for findings backed by an approved AFR document type."""
    if fmt not in {'csv', 'xlsx'}:
        raise ValueError('Supported export formats are csv and xlsx')
    rows = _rows(db, audit_session_id)
    created = datetime.now(timezone.utc)
    output = get_settings().output_dir / f'Evidence_Findings_session-{audit_session_id}_{created:%Y%m%dT%H%M%SZ}_{uuid4().hex[:8]}.{fmt}'
    if fmt == 'csv':
        content = StringIO(newline='')
        writer = csv.writer(content)
        writer.writerow(HEADERS)
        writer.writerows(rows)
        output.write_text(content.getvalue(), encoding='utf-8-sig', newline='')
    else:
        from openpyxl import Workbook
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Open Findings'
        sheet.append(HEADERS)
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = 'A2'
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(60, max(14, max(len(str(cell.value or '')) for cell in column) + 2))
        if not rows:
            add_empty_findings_message(sheet, len(HEADERS))
        detailed_findings = db.scalars(evidence_findings_query(audit_session_id).where(
            Finding.finding_kind != 'document_catalogue'
        ).order_by(Finding.severity, Finding.practice_area_code, Finding.id)).all()
        detailed_context = _detailed_control_context(db, audit_session_id)
        _append_findings_sheet(
            workbook,
            'Detailed Control Findings',
            [[*detailed_context, *row] for row in _rows(db, audit_session_id, detailed_findings)],
            DETAILED_CONTROL_HEADERS,
        )
        _append_classification_sheet(workbook, db, audit_session_id)
        _append_document_coverage_sheet(workbook, db, audit_session_id)
        scope_status, unreadable_paths = audit_scope_status(db, audit_session_id)
        metadata = workbook.create_sheet('Report Metadata')
        metadata.append(['Audit session ID', audit_session_id])
        metadata.append(['Evidence assessment scope', scope_status])
        metadata.append(['Unreadable evidence file(s)', '\n'.join(unreadable_paths) or 'None'])
        metadata.append(['Report scope', 'Open Findings contains confirmed file-backed findings and project-scope coverage gaps. Detailed Control Findings contains the file-backed auditor-review queue. Evidence Classification records the document mapping used for this scan.'])
        workbook.save(output)
    report = GeneratedReport(audit_session_id=audit_session_id, report_type=f'evidence_findings_{fmt}',
                             storage_path=str(output), created_by_id=user_id)
    db.add(report)
    db.commit()
    db.refresh(report)
    return report
