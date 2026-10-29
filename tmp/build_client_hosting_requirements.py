from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'deliverables' / 'CMMI_Audit_Platform_Client_Hosting_Installation_and_Infrastructure_Requirements.docx'
OUT.parent.mkdir(exist_ok=True)

NAVY = '17365D'
BLUE = 'DCE6F1'
PALE = 'F5F8FC'
GRID = 'D9D9D9'
GREY = RGBColor(89, 89, 89)


def remove_paragraph_border(paragraph_or_style):
    p_pr = paragraph_or_style._element.get_or_add_pPr()
    border = p_pr.find(qn('w:pBdr'))
    if border is not None:
        p_pr.remove(border)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_border(cell, color=GRID):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = qn(f'w:{edge}')
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f'w:{edge}')
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), color)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in('w:tcMar')
    if mar is None:
        mar = OxmlElement('w:tcMar')
        tc_pr.append(mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    hdr = OxmlElement('w:tblHeader')
    hdr.set(qn('w:val'), 'true')
    tr_pr.append(hdr)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    no_split = OxmlElement('w:cantSplit')
    tr_pr.append(no_split)


def set_column_widths(table, widths):
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    repeat_header(header)
    for cell, text in zip(header.cells, headers):
        cell.text = text
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.runs[0]
        r.font.name = 'Aptos'
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
    for index, values in enumerate(rows):
        added_row = table.add_row()
        keep_row_together(added_row)
        cells = added_row.cells
        for cell, value in zip(cells, values):
            cell.text = value
            set_cell_border(cell)
            set_cell_margins(cell)
            if index % 2:
                set_cell_shading(cell, PALE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                for r in p.runs:
                    r.font.name = 'Aptos'
                    r.font.size = Pt(8.5)
    set_column_widths(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style='Heading 1' if level == 1 else 'Heading 2')
    p.paragraph_format.keep_with_next = True
    p.add_run(text)
    return p


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def add_check(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(4)
    p.add_run('☐  ' + text)
    return p


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.68)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10)
styles['Normal'].font.color.rgb = RGBColor(0, 0, 0)
styles['Normal'].paragraph_format.space_after = Pt(7)
for name, size in [('Title', 20), ('Heading 1', 14), ('Heading 2', 11)]:
    styles[name].font.name = 'Aptos Display'
    styles[name]._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos Display')
    styles[name]._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos Display')
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
styles['Title'].paragraph_format.space_after = Pt(7)
styles['Heading 1'].paragraph_format.space_before = Pt(14)
styles['Heading 1'].paragraph_format.space_after = Pt(6)
styles['Heading 2'].paragraph_format.space_before = Pt(9)
styles['Heading 2'].paragraph_format.space_after = Pt(4)
remove_paragraph_border(styles['Title'])

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('CMMI Audit Platform Client Hosting Installation and Infrastructure Requirements')
remove_paragraph_border(title)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(16)
r = subtitle.add_run('Client prerequisites for hosting, operating and accessing the application')
r.font.size = Pt(10)
r.font.color.rgb = GREY

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.paragraph_format.space_after = Pt(15)
r = meta.add_run('Prepared for client IT, infrastructure and procurement review | 27 October 2026')
r.font.size = Pt(9)
r.font.color.rgb = GREY

add_heading(doc, 'A Document Overview')
add_para(doc, 'This document sets out the client-side prerequisites for hosting and accessing the CMMI Audit Platform. It distinguishes mandatory production requirements from recommended operational practices and optional integrations. The vendor will provide the application package, configuration template, licence notices and deployment instructions. The client or its nominated hosting provider will arrange the agreed production environment.')
add_para(doc, 'This document is suitable for client review once the confirmation items in Section I have been resolved. It does not commit either party to an implementation, support arrangement or cost unless that is agreed separately in writing.', 'Important: ')

add_heading(doc, 'B Minimum Server and Infrastructure Requirements')
add_para(doc, 'The current deployment design is a containerised application and MySQL database running on a Linux server or virtual machine. The supplied evidence does not establish capacity figures. CPU, memory and storage must therefore be sized through workload testing and the expected evidence, report and retention volumes.')
add_table(doc,
    ['Requirement', 'Minimum or recommended specification', 'Purpose', 'Status', 'Responsibility'],
    [
        ['Linux server or virtual machine', 'A supported Linux server or VM suitable for Docker Engine and Docker Compose. Exact distribution and version require technical confirmation.', 'Hosts the application, database and supporting containers.', 'Mandatory', 'Client hosting team'],
        ['CPU and RAM', 'To be confirmed through workload testing, expected concurrent users, document size and OCR workload.', 'Supports application processing, document conversion and database activity.', 'Mandatory sizing decision', 'Client with vendor input'],
        ['Persistent disk storage', 'Separate capacity for database, uploaded evidence, generated reports and templates. Retention and growth assumptions must be agreed.', 'Preserves business data across container upgrades.', 'Mandatory', 'Client hosting team'],
        ['Backup and recovery storage', 'Protected backup location for database and all persistent file volumes; restore testing arrangement required.', 'Supports recovery following failure, error or incident.', 'Mandatory', 'Client hosting team'],
        ['Network and firewall', 'Inbound HTTPS through the selected reverse proxy; application port remains restricted to the host or proxy. Database remains on a private network.', 'Provides controlled user access and protects internal services.', 'Mandatory', 'Client hosting team'],
        ['Domain name and DNS', 'A production hostname under client control or an agreed vendor-managed domain.', 'Provides the public application address and certificate validation.', 'Mandatory for external access', 'Client unless agreed otherwise'],
        ['HTTPS and TLS certificate', 'Valid certificate and HTTPS reverse-proxy configuration. Provider and renewal ownership to be agreed.', 'Encrypts browser-to-application traffic.', 'Mandatory in production', 'Client hosting team'],
        ['Monitoring and alerting', 'Operational monitoring, log retention and alert routing appropriate to the agreed support model.', 'Supports availability, incident response and capacity management.', 'Recommended; support-model dependent', 'Client or vendor by agreement'],
    ], [1.16, 2.35, 1.45, 0.78, 1.12])

add_heading(doc, 'C Software Required on the Server')
add_para(doc, 'The production server needs container runtime software and an approved deployment configuration. End-user computers do not host these components.')
add_table(doc,
    ['Component', 'Where provided or installed', 'Purpose and client action', 'Status'],
    [
        ['Docker Engine', 'Installed on the Linux production server by the client hosting team.', 'Runs the approved application and database containers.', 'Mandatory'],
        ['Docker Compose plugin', 'Installed on the Linux production server by the client hosting team.', 'Starts and manages the defined multi-container deployment.', 'Mandatory'],
        ['Application container image', 'Packaged and supplied by the vendor.', 'Runs the Python, NiceGUI, FastAPI and Uvicorn application components.', 'Mandatory'],
        ['MySQL Community container', 'Pulled from an approved registry or supplied by an agreed process.', 'Provides the application database and persistent data store.', 'Mandatory'],
        ['HTTPS reverse proxy', 'Selected and installed/configured by the client hosting team.', 'Terminates TLS and publishes the application securely. The specific product is not yet confirmed.', 'Mandatory in production'],
        ['LibreOffice Core and Tesseract OCR', 'Included in the current application-image design for document processing and OCR.', 'No separate client-server installation is expected if confirmed in the final image. Final image contents require verification.', 'Required for full document-processing capability'],
        ['Docker Desktop', 'Not required for a Linux production server.', 'A desktop development product; it is not the production-server prerequisite.', 'Not required'],
    ], [1.55, 1.95, 2.75, 1.06])

add_para(doc, 'The vendor will provide a configuration template. The client must provide approved production secrets, database credentials, hostname values and any selected third-party service credentials through the agreed secure process.', 'Configuration responsibility: ')

add_heading(doc, 'D Software Required on Client Computers')
add_para(doc, 'End users access the CMMI Audit Platform through a modern, supported web browser over HTTPS. Based on the documented browser-based architecture, users do not need to install Python, Docker, MySQL, LibreOffice, Tesseract, database tools or development and security-scanning tools on their computers.')
add_table(doc,
    ['Client computer item', 'Requirement', 'Remarks'],
    [
        ['Modern web browser', 'Mandatory', 'Use a currently supported browser and permit HTTPS access to the agreed application URL.'],
        ['Internet or corporate-network access', 'Mandatory', 'Network access must follow the client firewall and identity policies.'],
        ['Python, Docker, MySQL, LibreOffice and Tesseract', 'Not required', 'These are server-side or containerised components in the documented deployment design.'],
    ], [2.0, 1.25, 4.06])
add_para(doc, 'Development and release-assurance tools such as GitHub Actions, Bandit, Gitleaks, Trivy, pip-audit and pytest are not required on client computers. They remain part of the vendor development and release process.', 'Development tools: ')

add_heading(doc, 'E Software Licences and Subscriptions')
add_para(doc, 'The following position is based on the documented architecture and preliminary licence review. It is not a substitute for legal advice. The delivery package must include accurate third-party notices and an inventory of the final release image and dependencies.')
add_table(doc,
    ['Software or service', 'Licence or payment requirement', 'Who arranges it', 'Remarks'],
    [
        ['Python application stack', 'No separate commercial licence purchase identified for the documented permissive-licence packages; applicable notices and licence terms must be retained.', 'Vendor prepares notices; client receives them.', 'Includes the main Python web, database and document-processing libraries, subject to final SBOM verification.'],
        ['MySQL Community', 'No commercial licence purchase identified for the Community deployment path; GPL-related redistribution obligations depend on how the image is supplied.', 'Client or vendor according to agreed delivery model.', 'Client pull of an approved official image is the preferred path. MySQL Enterprise is a separate commercial product and is not required by the documented design.'],
        ['Docker Engine and Compose', 'No Docker Desktop subscription is required for a Linux production server.', 'Client hosting team.', 'Do not treat Docker Engine, Docker Compose and Docker Desktop as the same product or commercial requirement.'],
        ['LibreOffice Core and Tesseract', 'No separate commercial licence purchase identified for the documented server-side use; applicable open-source notices must accompany the final package.', 'Vendor prepares notices and image inventory.', 'Final operating-system package inventory remains to be verified.'],
        ['PyMuPDF', 'Unresolved material licensing condition. The installed component is dual licensed under AGPL or a commercial Artifex licence.', 'Vendor to resolve before delivery.', 'A commercial licence must be confirmed, or an approved compliant replacement/removal decision documented, if the proprietary scanned-PDF OCR capability is retained.'],
        ['Docker Desktop', 'Optional desktop product; paid subscription may be required depending on the organisation and use under Docker terms.', 'Organisation that elects to use it.', 'Not a client production-server requirement.'],
        ['Hosting, domain, DNS, backup storage and bandwidth', 'Provider charges depend on the selected environment and usage.', 'Client unless agreed otherwise.', 'Infrastructure operational expense, not an application software licence.'],
    ], [1.38, 2.35, 1.25, 2.33])

add_heading(doc, 'F External Services and Accounts')
add_table(doc,
    ['Service or account', 'Status', 'Purpose', 'Client action'],
    [
        ['Domain and DNS management', 'Required for externally accessible production deployment', 'Provides the application hostname and TLS validation path.', 'Provide or delegate management of the agreed domain and DNS records.'],
        ['TLS certificate provider', 'Required in production', 'Enables HTTPS.', 'Select provider and agree certificate issue/renewal ownership.'],
        ['SMTP or email-delivery account', 'Optional integration; required if password-reset emails are enabled', 'Sends account password-reset messages.', 'Decide whether the feature is required and provide approved credentials if enabled.'],
        ['Cloudflare Turnstile', 'Optional', 'Additional bot-protection control at sign-in when configured.', 'Decide whether to enable it; provide account and keys if selected.'],
        ['Microsoft Entra ID or OAuth', 'Not yet confirmed as implemented and selected', 'Potential enterprise authentication integration.', 'Do not procure or configure until technical implementation and client identity requirements are confirmed.'],
        ['Cloud hosting account', 'Optional deployment model', 'Provides VM, storage, network and managed-service resources.', 'Provide account or nominate the party responsible for hosting.'],
    ], [1.58, 1.54, 2.05, 2.14])

add_heading(doc, 'G Responsibilities Matrix')
add_para(doc, 'The allocation below describes the intended working model for planning purposes. It does not replace the parties’ final agreement.')
add_table(doc,
    ['Activity', 'Vendor development and deployment team', 'Client IT and hosting team'],
    [
        ['Application package and instructions', 'Provide the approved application package, configuration template, deployment instructions and supported upgrade process.', 'Review and apply the instructions within the agreed environment.'],
        ['Dependencies and notices', 'Document software dependencies and applicable third-party notices; resolve application packaging issues.', 'Retain supplied notices and operate the agreed environment in accordance with them.'],
        ['Hosting environment', 'Provide sizing input after workload details are available.', 'Provide or procure the server, storage, network connectivity and approved operating environment.'],
        ['Networking and HTTPS', 'State application connectivity requirements.', 'Configure firewall, DNS, reverse proxy, TLS certificate and access controls.'],
        ['Credentials and external services', 'Identify configuration fields and integration requirements.', 'Provide approved production credentials and selected external-service accounts through the agreed secure process.'],
        ['Backup, monitoring and infrastructure maintenance', 'Define application backup/upgrade considerations.', 'Arrange backups, recovery, monitoring, infrastructure patching and operational access under the agreed hosting model.'],
    ], [1.58, 2.65, 2.95])

add_heading(doc, 'H Client Prerequisites Checklist')
add_para(doc, 'The client should complete the following items before scheduling deployment:')
for item in [
    'A Linux server or virtual machine has been provisioned, with administrator access available to the authorised hosting team.',
    'CPU, memory and persistent-storage capacity have been agreed after expected users, document volumes and retention needs are reviewed.',
    'Docker Engine and Docker Compose are installed on the server, or an approved hosting team will install them.',
    'The production hostname, DNS ownership and HTTPS certificate approach have been agreed.',
    'Firewall rules allow HTTPS access through the reverse proxy and keep database access private.',
    'Persistent storage is available for MySQL data, uploads, generated reports and templates.',
    'Backup location, retention, encryption, access controls and restore-testing responsibility have been agreed.',
    'A designated client IT contact and escalation route have been identified.',
    'The client has decided whether SMTP, Turnstile and Microsoft authentication integrations are required.',
    'Production credentials and any chosen external-service accounts can be provided through an approved secure process.',
    'The PyMuPDF licensing decision and final container/SBOM evidence have been accepted before release delivery.',
]:
    add_check(doc, item)

add_heading(doc, 'I Assumptions and Items Requiring Confirmation')
add_table(doc,
    ['Item requiring confirmation', 'Why it matters', 'Required decision or evidence'],
    [
        ['Hardware sizing', 'No supported CPU, memory, throughput or storage figures are established by the supplied evidence.', 'Expected users, evidence volume, OCR workload, retention period and workload-test result.'],
        ['Final production image inventory', 'The full operating-system package and container dependency inventory has not yet been verified.', 'Pinned image digest, final SBOM, licence notices and vulnerability-scan record.'],
        ['PyMuPDF licence decision', 'The documented OCR functionality creates a material licensing condition for proprietary client delivery.', 'Commercial Artifex licence confirmation or approved compliant replacement/removal decision.'],
        ['Reverse proxy and certificate provider', 'The architecture requires HTTPS but does not select a product or service owner.', 'Selected product/provider and responsibility for configuration and renewal.'],
        ['SMTP, Turnstile and Microsoft authentication', 'These are optional or not fully confirmed integrations.', 'Client requirement, selected provider/account and configuration ownership.'],
        ['MySQL image delivery model', 'Redistribution obligations differ if the vendor supplies an image bundle rather than the client pulls an approved image.', 'Agreed delivery method and accompanying notice/source-compliance process.'],
        ['Support and operating model', 'Monitoring, backup, patching and incident responsibilities depend on the final agreement.', 'Named responsibilities, service boundaries and escalation process.'],
    ], [1.78, 2.55, 2.85])

add_para(doc, 'Once these confirmation items are resolved, this document can be issued with the final release package, configuration template, third-party notices, SBOM and deployment instructions.', 'Release readiness: ')

for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(4)
    run = footer.add_run('CMMI Audit Platform | Client Hosting Installation and Infrastructure Requirements')
    run.font.name = 'Aptos'
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(128, 128, 128)

doc.save(OUT)
print(OUT)
