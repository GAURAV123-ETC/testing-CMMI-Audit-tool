from io import BytesIO
from types import SimpleNamespace
import zipfile
import pytest

from sqlalchemy import create_engine, select
from sqlalchemy.dialects import mysql
from sqlalchemy.orm import sessionmaker
from sqlalchemy.schema import CreateTable

from app.db.database import Base
from app.db.models import (AuditSession, AuditSessionPracticeArea, EvidenceFile,
                           EvidenceSource, Finding, PracticeArea, ChecklistVersion,
                           Customer, AuditProject, CmmiRule, DocumentTypeRule, ScanJob, AuditLog)
from app.services.audit_engine import evidence_scan
from app.services import evidence_ingestion
from app.services.reports import afr_report, finding_export


def test_evidence_text_uses_longtext_on_mysql_for_large_workbooks():
    ddl = str(CreateTable(EvidenceFile.__table__).compile(dialect=mysql.dialect())).upper()
    assert 'EXTRACTED_TEXT LONGTEXT' in ddl


def test_stored_evidence_path_is_absolute_and_confined(monkeypatch, tmp_path):
    upload_dir = tmp_path / 'uploads'
    upload_dir.mkdir()
    monkeypatch.setattr(evidence_scan, 'get_settings', lambda: SimpleNamespace(upload_dir=upload_dir))

    resolved = evidence_scan._stored_evidence_path('uploads/14/evidence.doc')

    assert resolved == str((upload_dir / '14' / 'evidence.doc').resolve())
    with pytest.raises(ValueError, match='outside the configured upload directory'):
        evidence_scan._stored_evidence_path('../outside.doc')


def test_key_normalization_keeps_scalar_values_intact_and_removes_blank_duplicates():
    assert evidence_scan._unique_nonblank_keys('Closure date') == ['Closure date']
    assert evidence_scan._unique_nonblank_keys([' Closure date ', 'closure date', None, '']) == ['Closure date']


def add_test_catalog(db, version_id=1):
    db.add(ChecklistVersion(id=version_id, version=f'test-{version_id}', source='test', checksum=str(version_id) * 64,
                            framework='CMMI v3.0', status='ACTIVE', is_active=True))
    db.add_all([
        CmmiRule(checklist_version_id=version_id, rule_id='PLAN-G1', practice_area_code='PLAN', level='L1-Gate',
                 audit_check='Locate the approved project plan.', gap_text='Project plan missing.'),
        CmmiRule(checklist_version_id=version_id, rule_id='PLAN-1', practice_area_code='PLAN', level='L3-Check',
                 audit_check='Are schedule and milestone commitments approved?', gap_text='Planning commitments incomplete.'),
        CmmiRule(checklist_version_id=version_id, rule_id='IRP-G1', practice_area_code='IRP', level='L1-Gate',
                 audit_check='Locate the incident log.', gap_text='Incident log missing.'),
    ])
    db.add_all([
        DocumentTypeRule(checklist_version_id=version_id, document_type='Project Plan', practice_areas=['PLAN'],
                         keywords=['project plan', 'milestone', 'schedule', 'approved'], include_in_afr=True),
        DocumentTypeRule(checklist_version_id=version_id, document_type='Issue Log', practice_areas=['IRP'],
                         keywords=['issue id', 'priority', 'incident'], include_in_afr=True),
    ])


def test_persisted_evidence_scan_creates_findings_and_updates_practice_area_status(monkeypatch, tmp_path):
    engine = create_engine('sqlite://')
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine)
    upload_dir = tmp_path / 'uploads'
    stored_file = upload_dir / '1' / 'plan.md'
    stored_file.parent.mkdir(parents=True)
    stored_file.write_text('project plan milestone schedule and approved risk mitigation', encoding='utf-8')
    monkeypatch.setattr(evidence_scan, 'get_settings', lambda: SimpleNamespace(upload_dir=upload_dir))

    with session_factory() as db:
        add_test_catalog(db)
        db.add_all([PracticeArea(id=1, code='PLAN', name='Planning'), PracticeArea(id=2, code='IRP', name='Incident Resolution')])
        db.add(AuditSession(id=1, project_id=1, checklist_version_id=1, created_by_id=1))
        db.add_all([
            AuditSessionPracticeArea(audit_session_id=1, practice_area_id=1),
            AuditSessionPracticeArea(audit_session_id=1, practice_area_id=2),
        ])
        source = EvidenceSource(audit_session_id=1, source_type='test', created_by_id=1)
        db.add(source); db.flush()
        db.add(EvidenceFile(source_id=source.id, relative_path='Project/plan.md',
                            storage_path='uploads/1/plan.md', sha256='a' * 64, size_bytes=stored_file.stat().st_size))
        db.add(EvidenceFile(source_id=source.id, relative_path='Project/plan-copy.md',
                            storage_path='uploads/1/plan.md', sha256='a' * 64, size_bytes=stored_file.stat().st_size))
        db.commit()

        outcome = evidence_scan.scan_session(db, 1, 1)
        findings = db.scalars(select(Finding).where(Finding.audit_session_id == 1)).all()
        audit_actions = set(db.scalars(select(AuditLog.action)).all())
        statuses = {row.practice_area_id: (row.folder_status, row.evidence_status)
                    for row in db.scalars(select(AuditSessionPracticeArea).where(AuditSessionPracticeArea.audit_session_id == 1)).all()}

    assert outcome['files_processed'] == 1
    assert outcome['duplicate_files_skipped'] == 1
    assert outcome['findings_created'] == len(findings)
    assert outcome['findings_created'] == outcome['evidence_findings_created'] + outcome['coverage_gaps_created']
    assert all(isinstance(finding.required_keys, list) for finding in findings)
    assert all(isinstance(finding.available_keys, list) for finding in findings)
    assert all(isinstance(finding.missing_required_keys, list) for finding in findings)
    assert {'evidence_scan_started', 'evidence_scan_completed'} <= audit_actions
    assert statuses[1][0] == 'available'
    assert statuses[2] == ('not_scanned', 'not_scanned')
    assert not any(finding.rule_id.startswith('IRP-') for finding in findings)
    coverage = next(finding for finding in findings if finding.rule_id == 'DOC-COVERAGE-2')
    assert coverage.finding_kind == 'coverage_gap'
    assert coverage.required_keys == ['Issue Log']
    assert coverage.available_keys == []
    assert coverage.missing_required_keys == ['No uploaded file was classified as Issue Log.']
    assert outcome['coverage_gaps_created'] == 1
    assert db.get(EvidenceFile, 1).classification_json['detected_type'] == 'Project Plan'
    assert db.get(EvidenceFile, 1).classified_at is not None


def test_scan_uses_catalogue_scope_and_does_not_filename_trigger_irp_validation(monkeypatch, tmp_path):
    engine = create_engine('sqlite://')
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine)
    upload_dir = tmp_path / 'uploads'
    stored_file = upload_dir / '1' / 'issue-log.csv'
    stored_file.parent.mkdir(parents=True)
    stored_file.write_text('Issue ID,Priority\nINC-1,P1\n', encoding='utf-8')
    monkeypatch.setattr(evidence_scan, 'get_settings', lambda: SimpleNamespace(upload_dir=upload_dir))

    with session_factory() as db:
        add_test_catalog(db)
        issue_type = db.scalar(select(DocumentTypeRule).where(DocumentTypeRule.document_type == 'Issue Log'))
        issue_type.practice_areas = ['MC']
        db.add_all([PracticeArea(id=1, code='IRP', name='Incident Resolution'),
                    PracticeArea(id=2, code='MC', name='Manage Compliance')])
        db.add(AuditSession(id=1, project_id=1, checklist_version_id=1, created_by_id=1))
        db.add_all([AuditSessionPracticeArea(audit_session_id=1, practice_area_id=1),
                    AuditSessionPracticeArea(audit_session_id=1, practice_area_id=2)])
        source = EvidenceSource(audit_session_id=1, source_type='test', created_by_id=1)
        db.add(source); db.flush()
        db.add(EvidenceFile(source_id=source.id, relative_path='IRP/Issue Log.csv',
                            storage_path='uploads/1/issue-log.csv', sha256='b' * 64,
                            size_bytes=stored_file.stat().st_size))
        db.commit()

        outcome = evidence_scan.scan_session(db, 1, 1)
        rule_ids = set(db.scalars(select(Finding.rule_id).where(Finding.audit_session_id == 1)).all())

    assert outcome['findings_created'] > 0
    assert 'IRP-HEADER' not in rule_ids
    assert db.get(EvidenceFile, 1).classification_json['practice_areas'] == ['MC']


def test_document_keys_flow_unchanged_from_spreadsheet_to_persisted_finding_and_afr(monkeypatch, tmp_path):
    """Exercise extraction, table matching, persistence, and the final XLSX row."""
    engine = create_engine('sqlite://')
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine)
    upload_dir = tmp_path / 'uploads'
    stored_file = upload_dir / '1' / 'change-log.csv'
    stored_file.parent.mkdir(parents=True)
    stored_file.write_text(
        'Change ID,Closure Date\nCR-1,2026-09-20\nCR-2,\n', encoding='utf-8'
    )
    monkeypatch.setattr(evidence_scan, 'get_settings', lambda: SimpleNamespace(upload_dir=upload_dir))
    monkeypatch.setattr(afr_report, 'get_settings', lambda: SimpleNamespace(output_dir=tmp_path))
    monkeypatch.setattr(finding_export, 'get_settings', lambda: SimpleNamespace(output_dir=tmp_path))

    with session_factory() as db:
        db.add(ChecklistVersion(id=1, version='test', source='test', checksum='e' * 64,
                                framework='CMMI v3.0', status='ACTIVE', is_active=True))
        document_type = DocumentTypeRule(
            checklist_version_id=1, document_type='Change Log', practice_areas=['RDM'],
            keywords=['Change ID', 'Closure Date'], include_in_afr=True,
        )
        db.add(document_type)
        db.flush()
        db.add_all([
            CmmiRule(
                checklist_version_id=1, document_type_rule_id=document_type.id,
                rule_id='CR-G1', practice_area_code='RDM', level='L1-Gate',
                audit_check='Locate the approved change log.', gap_text='Change log missing.',
            ),
            CmmiRule(
                checklist_version_id=1, document_type_rule_id=document_type.id,
                rule_id='CR-19', practice_area_code='RDM', level='L3-Check',
                detection='Closure date', audit_check='Closure date is recorded for each change.',
                gap_text='Closure date is missing.', recommendation='Record a closure date for each change.',
            ),
        ])
        db.add(PracticeArea(id=1, code='RDM', name='Requirements Development and Management'))
        db.add(AuditSession(id=1, project_id=1, checklist_version_id=1, created_by_id=1))
        db.add(AuditSessionPracticeArea(audit_session_id=1, practice_area_id=1))
        source = EvidenceSource(audit_session_id=1, source_type='test', created_by_id=1)
        db.add(source)
        db.flush()
        db.add(EvidenceFile(
            source_id=source.id, relative_path='RDM/change-log.csv',
            storage_path='uploads/1/change-log.csv', sha256='f' * 64,
            size_bytes=stored_file.stat().st_size,
        ))
        db.commit()

        evidence_scan.scan_session(db, 1, 1)
        finding = db.scalars(select(Finding).where(
            Finding.audit_session_id == 1, Finding.rule_id == 'CR-19', Finding.status == 'open',
        )).one()
        assert finding.required_keys == ['Closure date']
        assert finding.available_keys == ['Closure Date: 1 populated value(s)']
        assert finding.missing_required_keys == ['Closure Date: 1 blank value(s)']

        report = afr_report.create_afr(db, 1, 1, 'xlsx')
        report_path = report.storage_path
        evidence_scan_export = finding_export.create_current_findings_export(db, 1, 1, 'xlsx')
        evidence_scan_export_path = evidence_scan_export.storage_path

    from openpyxl import load_workbook
    workbook = load_workbook(report_path, read_only=True)
    headers = next(workbook['Detailed Findings'].iter_rows(values_only=True))
    values = dict(zip(headers, next(workbook['Detailed Findings'].iter_rows(min_row=2, values_only=True))))
    assert values['Required Keys'] == 'Closure date'
    assert values['Available Keys'] == 'Closure Date: 1 populated value(s)'
    assert values['Missing Required Keys'] == 'Closure Date: 1 blank value(s)'
    assert values['Finding'] == 'Partial evidence for CR-19'
    export_workbook = load_workbook(evidence_scan_export_path, read_only=True)
    export_headers = next(export_workbook['Detailed Control Findings'].iter_rows(values_only=True))
    export_values = dict(zip(
        export_headers,
        next(export_workbook['Detailed Control Findings'].iter_rows(min_row=2, values_only=True)),
    ))
    assert export_values['Required Keys'] == values['Required Keys']
    assert export_values['Available Keys'] == values['Available Keys']
    assert export_values['Missing Required Keys'] == values['Missing Required Keys']


def test_scan_rejects_a_second_running_job(tmp_path):
    engine = create_engine('sqlite://')
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine)
    with session_factory() as db:
        add_test_catalog(db)
        db.add(AuditSession(id=1, project_id=1, checklist_version_id=1, created_by_id=1))
        db.add(ScanJob(audit_session_id=1, job_type='master_rule_and_irp_scan', status='running'))
        db.commit()
        try:
            evidence_scan.scan_session(db, 1, 1)
        except ValueError as exc:
            assert 'already running' in str(exc)
        else:
            raise AssertionError('Expected an existing running scan to be rejected.')


def test_failed_scan_records_a_user_actionable_job_error():
    engine = create_engine('sqlite://')
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine)
    with session_factory() as db:
        add_test_catalog(db)
        db.add(AuditSession(id=1, project_id=1, checklist_version_id=1, created_by_id=1))
        db.commit()
        with pytest.raises(ValueError, match='No evidence files'):
            evidence_scan.scan_session(db, 1, 1)
        job = db.scalars(select(ScanJob).where(ScanJob.audit_session_id == 1)).one()

    assert job.status == 'failed'
    assert job.result_summary['error'] == 'No evidence files have been uploaded for this project.'


def test_parent_folder_archive_creates_one_persisted_audit_session_per_project(monkeypatch, tmp_path):
    engine = create_engine('sqlite://')
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine)
    upload_dir = tmp_path / 'uploads'
    monkeypatch.setattr(evidence_ingestion, 'get_settings', lambda: SimpleNamespace(upload_dir=upload_dir, max_upload_mb=50))
    archive = BytesIO()
    with zipfile.ZipFile(archive, 'w') as zipped:
        zipped.writestr('Project Alpha/plan.md', 'planning schedule milestone')
        zipped.writestr('Project Beta/risk.md', 'risk impact owner')

    with session_factory() as db:
        db.add(Customer(id=1, name='Customer A', created_by_id=1))
        db.add(ChecklistVersion(id=1, version='test', source='test', checksum='c' * 64, is_active=True))
        db.add_all([PracticeArea(id=1, code='PLAN', name='Planning'), PracticeArea(id=2, code='RSK', name='Risk')])
        db.commit()
        created = evidence_ingestion.import_parent_folder_archive(
            db, 1, 1, 'parent-projects.zip', archive.getvalue(), 'application/zip'
        )
        db.commit()
        audits = db.scalars(select(AuditSession).order_by(AuditSession.id)).all()
        projects = db.scalars(select(AuditProject).order_by(AuditProject.name)).all()
        relative_paths = db.scalars(select(EvidenceFile.relative_path).order_by(EvidenceFile.relative_path)).all()

    assert [(item['project_name'], item['files_imported']) for item in created] == [('Project Alpha', 1), ('Project Beta', 1)]
    assert len(audits) == 2
    assert [project.name for project in projects] == ['Project Alpha', 'Project Beta']
    assert relative_paths == ['plan.md', 'risk.md']


def test_bulk_import_rejects_a_single_project_zip(monkeypatch, tmp_path):
    engine = create_engine('sqlite://')
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine)
    monkeypatch.setattr(evidence_ingestion, 'get_settings', lambda: SimpleNamespace(upload_dir=tmp_path / 'uploads', max_upload_mb=50))
    archive = BytesIO()
    with zipfile.ZipFile(archive, 'w') as zipped:
        zipped.writestr('Project Alpha/PLAN/plan.md', 'planning schedule milestone')

    with session_factory() as db:
        db.add(Customer(id=1, name='Customer A', created_by_id=1))
        db.add(ChecklistVersion(id=1, version='test', source='test', checksum='d' * 64, is_active=True))
        db.add(PracticeArea(id=1, code='PLAN', name='Planning'))
        db.commit()
        with pytest.raises(ValueError, match='at least two top-level project folders'):
            evidence_ingestion.import_parent_folder_archive(db, 1, 1, 'one-project.zip', archive.getvalue())
