from datetime import date
from types import SimpleNamespace
from pathlib import Path

from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db.database import Base
from app.db.models import (AuditProject, AuditSession, Comment, DocumentTypeRule, EvidenceFile,
                           EvidenceSource, Finding, FindingEvidence, RemediationAction)
from app.services.reports import afr_report, finding_export
from app.services.finding_scope import afr_findings_query


def test_key_value_rendering_handles_legacy_scalar_and_non_string_values():
    assert finding_export.format_key_values('Closure date') == 'Closure date'
    assert finding_export.format_key_values(['Closure date', 7, None, '']) == 'Closure date\n7'
    assert finding_export.format_available_keys([], has_linked_evidence=True) == (
        'No reliably matched available key was found in this document.'
    )
    assert finding_export.format_available_keys(None, has_linked_evidence=False) == (
        'No matching document or available key was found in the audit scope.'
    )
    assert finding_export.display_rule_id('DOC-COVERAGE-57') == 'Document Coverage'
    assert finding_export.display_rule_id('FSD-12') == 'FSD-12'


def test_empty_findings_sheet_explains_that_the_export_succeeded():
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(finding_export.HEADERS)
    finding_export.add_empty_findings_message(sheet, len(finding_export.HEADERS))

    assert sheet['A2'].value == finding_export.NO_CONFIRMED_FINDINGS
    assert 'A2:N2' in {str(cell_range) for cell_range in sheet.merged_cells.ranges}


def test_reference_document_classification_is_explicit_not_blank():
    row = finding_export.classification_display_values('Alpha/Manifest.xlsx', {
        'detected_type': 'Manifest / Mapping Reference', 'confidence': 'High',
        'classification_reason': 'Reference mapping only.', 'document_role': 'PROCESS_REFERENCE',
    })
    assert row[2] == 'Not applicable - reference document'
    assert row[4] == 'Not applicable - reference document'
    assert row[5] == 'Not applicable - reference document'


def test_implementation_classification_uses_explicit_empty_key_messages():
    row = finding_export.classification_display_values('Alpha/Unknown.xlsx', {
        'document_role': 'PROJECT_IMPLEMENTATION_EVIDENCE',
    })
    assert row[2] == 'Not classified'
    assert row[4] == 'No configured document keys matched'
    assert row[5] == 'No missing classification keys identified'


def test_document_coverage_uses_persisted_scan_classification_and_lists_missing_types():
    engine = create_engine('sqlite://')
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    from openpyxl import Workbook

    with factory() as db:
        db.add(AuditSession(id=1, project_id=1, checklist_version_id=7, created_by_id=1))
        matched_type = DocumentTypeRule(
            id=10, checklist_version_id=7, document_type='Estimation Sheet',
            practice_areas=['EST'], keywords=['estimate id'], aliases=[],
            expected_evidence='Approved estimation record', include_in_afr=True,
        )
        missing_type = DocumentTypeRule(
            id=11, checklist_version_id=7, document_type='WBS',
            practice_areas=['PLAN'], keywords=['wbs'], aliases=[],
            expected_evidence='Work breakdown structure', include_in_afr=True,
        )
        source = EvidenceSource(audit_session_id=1, source_type='test', created_by_id=1)
        db.add_all([matched_type, missing_type, source]); db.flush()
        db.add(EvidenceFile(
            source_id=source.id, relative_path='Alpha/Estimation.xlsx', storage_path='uploads/est.xlsx',
            sha256='f' * 64, size_bytes=1, processing_status='processed', classification_json={
                'document_type_rule_id': 10, 'document_role': 'PROJECT_IMPLEMENTATION_EVIDENCE',
                'detected_type': 'Estimation Sheet', 'confidence': 'High',
            },
        ))
        db.commit()
        workbook = Workbook()
        workbook.remove(workbook.active)
        finding_export._append_document_coverage_sheet(workbook, db, 1)

    rows = list(workbook['Document Coverage'].iter_rows(values_only=True))
    values = {row[0]: row for row in rows[1:]}
    assert values['Estimation Sheet'][3] == 'Matched'
    assert values['Estimation Sheet'][4] == 'Alpha/Estimation.xlsx'
    assert values['WBS'][3] == 'Not matched in uploaded evidence'
    assert values['WBS'][4] == 'No uploaded file matched this checklist document type.'


def test_afr_summary_includes_confirmed_sap_control_findings_not_review_queue():
    engine = create_engine('sqlite://')
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    with factory() as db:
        db.add(AuditSession(id=1, project_id=1, checklist_version_id=1, created_by_id=1))
        db.flush()
        source = EvidenceSource(audit_session_id=1, source_type='test', created_by_id=1)
        db.add(source); db.flush()
        evidence = EvidenceFile(source_id=source.id, relative_path='Project/evidence.xlsx',
                                storage_path='uploads/evidence.xlsx', sha256='a' * 64, size_bytes=1,
                                classification_json={'afr_eligible': True})
        db.add(evidence); db.flush()
        confirmed = Finding(audit_session_id=1, rule_id='CR-01', finding_kind='rule_assessment',
                            document_type_rule_id=1, practice_area_code='RDM', severity='major',
                            title='Missing evidence for CR-01', description='ID duplicate.',
                            recommendation='Correct it.', status='open')
        review = Finding(audit_session_id=1, rule_id='CR-03', finding_kind='rule_assessment',
                         document_type_rule_id=1, practice_area_code='RDM', severity='minor',
                         title='Review_Required evidence for CR-03', description='Confirm field.',
                         recommendation='Review it.', status='open')
        coverage = Finding(audit_session_id=1, rule_id='CR-GAP', finding_kind='coverage_gap',
                           practice_area_code='RDM', severity='major',
                           title='Missing evidence for CR-GAP', description='No matching artefact.',
                           recommendation='Upload the required artefact.', status='open')
        db.add_all([confirmed, review, coverage]); db.flush()
        db.add_all([FindingEvidence(finding_id=confirmed.id, evidence_file_id=evidence.id),
                    FindingEvidence(finding_id=review.id, evidence_file_id=evidence.id)])
        db.commit()
        result = db.scalars(afr_findings_query(1)).all()
    assert [finding.rule_id for finding in result] == ['CR-01', 'CR-GAP']


def test_current_findings_export_contains_only_open_selected_session_findings(monkeypatch, tmp_path):
    engine = create_engine('sqlite://')
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    monkeypatch.setattr(finding_export, 'get_settings', lambda: SimpleNamespace(output_dir=tmp_path))
    with factory() as db:
        db.add_all([
            AuditProject(id=1, customer_id=1, name='Acme Payments', owner_id=1),
            AuditSession(id=1, project_id=1, checklist_version_id=1, created_by_id=1,
                         audit_date=date(2026, 9, 21), auditors='Asha, Ravi', auditees='Payments team'),
            AuditSession(id=2, project_id=2, checklist_version_id=1, created_by_id=1),
            Finding(audit_session_id=1, rule_id='PLAN-G1', practice_area_code='PLAN', severity='major',
                    finding_kind='document_catalogue', document_type_rule_id=1,
                    title='=Plan missing', description='No approved plan.', recommendation='Upload the plan.', status='open',
                    required_keys=['schedule', 'approval'], available_keys=['schedule'], missing_required_keys=['approval']),
            Finding(audit_session_id=1, rule_id='PLAN-2', practice_area_code='PLAN', severity='minor',
                    title='Old result', description='Superseded.', recommendation='Ignore.', status='superseded'),
            Finding(audit_session_id=1, rule_id='BC-G1', practice_area_code='BC', severity='major',
                    title='BC plan not matched', description='Coverage gap.', recommendation='Review scope.', status='open'),
            Finding(audit_session_id=1, rule_id='XSESSION-G1', practice_area_code='PLAN', severity='major',
                    title='Foreign-session evidence link', description='Must not be treated as direct evidence.', recommendation='Review scope.', status='open'),
            Finding(audit_session_id=2, rule_id='RSK-G1', practice_area_code='RSK', severity='major',
                    title='Other session', description='Other.', recommendation='Other.', status='open'),
        ])
        db.flush()
        source = EvidenceSource(audit_session_id=1, source_type='test', created_by_id=1)
        db.add(source); db.flush()
        evidence = EvidenceFile(source_id=source.id, relative_path='ABC.zip/Project/approved-plan.docx',
                                storage_path='uploads/approved-plan.docx', sha256='a' * 64, size_bytes=1,
                                classification_json={'detected_type': 'Project Plan', 'practice_areas': ['PLAN'], 'afr_eligible': True})
        unreadable = EvidenceFile(source_id=source.id, relative_path='ABC.zip/Project/legacy-plan.doc',
                                  storage_path='uploads/legacy-plan.doc', sha256='c' * 64, size_bytes=1,
                                  processing_status='error')
        db.add_all([evidence, unreadable]); db.flush()
        db.add(FindingEvidence(finding_id=1, evidence_file_id=evidence.id))
        parse_error = Finding(audit_session_id=1, rule_id='FILE-PARSE', practice_area_code='UNCLASSIFIED', severity='major',
                              title='Could not parse legacy-plan.doc', description='Unreadable evidence.',
                              recommendation='Replace the file.', status='open')
        db.add(parse_error); db.flush()
        db.add(FindingEvidence(finding_id=parse_error.id, evidence_file_id=unreadable.id))
        ui_only = EvidenceFile(source_id=source.id, relative_path='ABC.zip/Project/legacy-plan.docx',
                               storage_path='uploads/legacy-plan.docx', sha256='e' * 64, size_bytes=1,
                               classification_json={'detected_type': 'Project Plan', 'practice_areas': ['PLAN'], 'afr_eligible': False})
        db.add(ui_only); db.flush()
        mixed_scope = Finding(audit_session_id=1, rule_id='MIXED-SCOPE', practice_area_code='PLAN', severity='major',
                              finding_kind='document_catalogue', document_type_rule_id=2,
                              title='Mixed evidence scope', description='Stored direct finding.',
                              recommendation='Review scope.', status='open')
        db.add(mixed_scope); db.flush()
        db.add_all([
            FindingEvidence(finding_id=mixed_scope.id, evidence_file_id=evidence.id),
            FindingEvidence(finding_id=mixed_scope.id, evidence_file_id=ui_only.id),
        ])
        foreign_source = EvidenceSource(audit_session_id=2, source_type='test', created_by_id=1)
        db.add(foreign_source); db.flush()
        foreign_evidence = EvidenceFile(source_id=foreign_source.id, relative_path='Other/foreign-plan.docx',
                                        storage_path='uploads/foreign-plan.docx', sha256='d' * 64, size_bytes=1)
        db.add(foreign_evidence); db.flush()
        # Simulate malformed legacy data: the association FKs are valid, but
        # the file belongs to another audit session.
        db.add_all([
            FindingEvidence(finding_id=1, evidence_file_id=foreign_evidence.id),
            FindingEvidence(finding_id=4, evidence_file_id=foreign_evidence.id),
        ])
        db.commit()
        csv_report = finding_export.create_current_findings_export(db, 1, 1, 'csv')
        xlsx_report = finding_export.create_current_findings_export(db, 1, 1, 'xlsx')
        csv_path = csv_report.storage_path
        xlsx_path = xlsx_report.storage_path

    csv_content = Path(csv_path).read_text(encoding='utf-8-sig')
    assert 'PLAN-G1' in csv_content
    assert "'=Plan missing" in csv_content
    assert 'PLAN-2' not in csv_content
    assert 'BC-G1' not in csv_content
    assert 'XSESSION-G1' not in csv_content
    assert 'RSK-G1' not in csv_content
    assert 'FILE-PARSE' not in csv_content
    assert 'MIXED-SCOPE' in csv_content
    assert 'Evidence Link Status' in csv_content
    assert 'Matched Document Type(s)' not in csv_content
    assert 'Required Keys' in csv_content
    assert 'schedule' in csv_content
    assert 'approval' in csv_content
    assert 'ABC.zip/Project/approved-plan.docx' in csv_content
    assert 'ABC.zip/Project/legacy-plan.docx' not in csv_content
    assert 'Other/foreign-plan.docx' not in csv_content
    assert 'assessment scope incomplete (1 unreadable file(s))' in csv_content
    assert xlsx_path.endswith('.xlsx')
    from openpyxl import load_workbook
    workbook = load_workbook(xlsx_path, read_only=True)
    assert workbook.sheetnames == ['Open Findings', 'Detailed Control Findings', 'Evidence Classification', 'Document Coverage', 'Report Metadata']
    coverage_rows = list(workbook['Document Coverage'].iter_rows(values_only=True))
    assert coverage_rows[0] == finding_export.DOCUMENT_COVERAGE_HEADERS
    detail_headers = next(workbook['Detailed Control Findings'].iter_rows(values_only=True))
    detail_rows = list(workbook['Detailed Control Findings'].iter_rows(min_row=2, values_only=True))
    assert detail_headers == finding_export.DETAILED_CONTROL_HEADERS
    detail_values = dict(zip(detail_headers, detail_rows[0]))
    assert detail_values['Project Name'] == 'Acme Payments'
    assert detail_values['Audit Date'] == '2026-09-21'
    assert detail_values['Auditor(s)'] == 'Asha, Ravi'
    assert detail_values['Auditee(s)'] == 'Payments team'
    assert any(row[5] == 'FILE-PARSE' for row in detail_rows)


def test_afr_bulk_loads_comments_and_remediation_actions(monkeypatch, tmp_path):
    engine = create_engine('sqlite://')
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine)
    monkeypatch.setattr(afr_report, 'get_settings', lambda: SimpleNamespace(output_dir=tmp_path))
    with factory() as db:
        db.add(AuditProject(id=1, customer_id=1, name='Gamma', owner_id=1))
        db.add(AuditSession(
            id=1, project_id=1, checklist_version_id=1, created_by_id=1,
            audit_date=date(2026, 9, 17), auditors='Asha, Ravi', auditees='Gamma delivery team',
        ))
        db.flush()
        finding = Finding(
            audit_session_id=1,
            rule_id='PLAN-G1',
            finding_kind='document_catalogue',
            document_type_rule_id=1,
            practice_area_code='PLAN',
            severity='major',
            title='Plan missing',
            description='No approved plan.',
            recommendation='Upload the plan.',
            status='open',
            required_keys=['schedule', 'approval'],
            available_keys=['schedule'],
            missing_required_keys=['approval'],
        )
        coverage_gap = Finding(
            audit_session_id=1, rule_id='BC-G1', practice_area_code='BC', severity='major',
            title='BC plan not matched', description='Coverage gap.', recommendation='Review scope.', status='open',
        )
        superseded = Finding(
            audit_session_id=1, rule_id='PLAN-OLD', practice_area_code='PLAN', severity='minor',
            title='Superseded direct finding', description='Historical.', recommendation='Ignore.', status='superseded',
        )
        control_finding = Finding(
            audit_session_id=1, rule_id='PLAN-C2', finding_kind='rule_assessment', document_type_rule_id=1,
            practice_area_code='PLAN', severity='minor', title='Review_Required evidence for PLAN-C2',
            description='Review the field-level evidence.', recommendation='Confirm the evidence.', status='open',
        )
        db.add_all([finding, coverage_gap, superseded, control_finding])
        db.flush()
        db.add_all([
            Comment(finding_id=finding.id, author_id=1, body='Reviewed by auditor.'),
            RemediationAction(finding_id=finding.id, action='Publish the approved plan.', status='open'),
        ])
        source = EvidenceSource(audit_session_id=1, source_type='test', created_by_id=1)
        db.add(source); db.flush()
        evidence = EvidenceFile(source_id=source.id, relative_path='ABC.zip/Project/approved-plan.docx',
                                storage_path='uploads/approved-plan.docx', sha256='b' * 64, size_bytes=1,
                                classification_json={'detected_type': 'Project Plan', 'practice_areas': ['PLAN'], 'afr_eligible': True})
        db.add(evidence); db.flush()
        db.add(FindingEvidence(finding_id=finding.id, evidence_file_id=evidence.id))
        db.add(FindingEvidence(finding_id=superseded.id, evidence_file_id=evidence.id))
        db.add(FindingEvidence(finding_id=control_finding.id, evidence_file_id=evidence.id))
        db.commit()
        report = afr_report.create_afr(db, 1, 1, 'xlsx')
        xlsx_path = report.storage_path
        docx_report = afr_report.create_afr(db, 1, 1, 'docx')
        docx_path = docx_report.storage_path
        pdf_report = afr_report.create_afr(db, 1, 1, 'pdf')
        pdf_path = pdf_report.storage_path

    from openpyxl import load_workbook
    worksheet = load_workbook(xlsx_path, read_only=True)['Detailed Findings']
    headers = next(worksheet.iter_rows(values_only=True))
    row = next(worksheet.iter_rows(min_row=2, values_only=True))
    values = dict(zip(headers, row))
    assert values['Evidence Link Status'] == 'Uploaded evidence assessed for this finding'
    assert values['Evidence File(s)'] == 'ABC.zip/Project/approved-plan.docx'
    assert values['Required Keys'] == 'schedule\napproval'
    assert values['Available Keys'] == 'schedule'
    assert values['Missing Required Keys'] == 'approval'
    assert values['Comments'] == 'Reviewed by auditor.'
    assert values['Remediation actions'] == 'open: Publish the approved plan.'
    assert all(row[1] != 'BC-G1' for row in worksheet.iter_rows(min_row=2, values_only=True))
    assert all(row[1] != 'PLAN-OLD' for row in worksheet.iter_rows(min_row=2, values_only=True))
    styled_worksheet = load_workbook(xlsx_path)['Detailed Findings']
    assert styled_worksheet.freeze_panes == 'A2'
    assert styled_worksheet.auto_filter.ref == 'A1:O2'
    assert styled_worksheet['A1'].font.bold is True
    assert styled_worksheet['A2'].alignment.wrap_text is True
    workbook = load_workbook(xlsx_path, read_only=True)
    assert workbook.sheetnames == ['Detailed Findings', 'Detailed Control Findings', 'Evidence Classification', 'Document Coverage', 'Report Metadata']
    control_sheet = workbook['Detailed Control Findings']
    control_headers = next(control_sheet.iter_rows(values_only=True))
    assert control_headers[:4] == finding_export.DETAILED_CONTROL_CONTEXT_HEADERS
    control_values = dict(zip(control_headers, next(control_sheet.iter_rows(min_row=2, values_only=True))))
    assert control_values['Project Name'] == 'Gamma'
    assert control_values['Audit Date'] == '2026-09-17'
    assert control_values['Auditor(s)'] == 'Asha, Ravi'
    assert control_values['Auditee(s)'] == 'Gamma delivery team'
    assert control_values['Available Keys'] == 'No reliably matched available key was found in this document.'
    assert load_workbook(xlsx_path)['Detailed Control Findings'].auto_filter.ref == 'A1:S2'
    assert Path(docx_path).is_file()
    assert Path(pdf_path).is_file()
