from io import BytesIO
import zipfile
import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker

from app.db.database import Base
from app.db.models import ChecklistVersion, CmmiRule, DocumentTypeRule, PracticeArea
from app.services.rule_catalog_importer import (
    _rule_signature,
    activate_rule_catalog,
    import_document_catalogue_additions,
    import_rule_catalog,
)


def test_rule_signature_includes_sap_executable_control_fields():
    base = {
        'practice_area': 'CAR', 'level': 'L3-Check', 'audit_check': 'RCA ID is present.',
        'gap_text': 'RCA ID missing.', 'document_key': 'rcaforp1p2', 'condition': 'Always',
        'detection': 'Unique RCA ID', 'recommendation': 'Assign an RCA ID.', 'is_mandatory': True,
    }
    revised = {**base, 'recommendation': 'Assign a unique RCA ID.'}
    assert _rule_signature(base) != _rule_signature(revised)


def workbook_bytes(duplicate=False, audit_check='Locate Risk Register.', duplicate_rule_id='RSK-G1'):
    book=Workbook(); master=book.active; master.title='MASTER'
    master.append(['#','Level','Rule ID','Gate / Artefact','Audit Check (What the auditor looks for and verifies)','Conformance','If NO: Finding and Implication'])
    master.append([1,'L1-Gate','RSK-G1','Artefact',audit_check,'', 'Risk Register missing.'])
    if duplicate: master.append([2,'L3-Check',duplicate_rule_id,'', 'Check risks.','', 'Risks incomplete.'])
    s2=book.create_sheet('Sheet2'); s2.append(['S. No.','Domain Name','Practice Area Name','Related Document Name','Main Keywords','Example Evidence / Example Document']); s2.append([1,'Development','Risk Management (RSK)','Risk Register','Risk ID, Probability, Impact, Owner, Mitigation','Current risk register'])
    s5=book.create_sheet('Sheet5'); s5.append(['Sr. No.','Document / Evidence','Related Practice Area','What to Check / Key Points','Expected Evidence']); s5.append([1,'Risk Register','RSK','Risk ID, Probability, Impact, Owner','Updated risk register'])
    s6=book.create_sheet('Sheet6'); s6.append(['Document / Evidence','Primary Purpose']); s6.append(['Risk Register','Manage risks'])
    output=BytesIO(); book.save(output); return output.getvalue()


def document_additions_bytes(rows):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Additional Evidence'
    sheet.append(['S.No', 'Document / Evidence', 'Related Practice Area',
                  'What to Check / Key Points', 'Expected Evidence'])
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def test_import_validates_then_explicitly_activates_immutable_catalog():
    engine=create_engine('sqlite://'); Base.metadata.create_all(engine); factory=sessionmaker(bind=engine)
    with factory() as db:
        result=import_rule_catalog(db,'master.xlsx',workbook_bytes()); db.commit()
        version=db.get(ChecklistVersion,result['checklist_version_id'])
        documents = db.scalars(select(DocumentTypeRule)).all()
        assert version.status == 'VALIDATED' and not version.is_active
        assert db.scalar(select(CmmiRule.rule_id)) == 'RSK-G1'
        assert len(documents) == 14
        assert {document.document_type for document in documents} >= {'WBS', 'Impact Analysis and TSD'}
        assert 'Risk Register' not in {document.document_type for document in documents}
        assert all(document.include_in_afr for document in documents)
        activate_rule_catalog(db, version.id); db.commit()
        assert db.get(ChecklistVersion, version.id).status == 'ACTIVE'


def test_small_text_change_creates_a_validated_version_without_replacing_active_rules():
    engine=create_engine('sqlite://'); Base.metadata.create_all(engine); factory=sessionmaker(bind=engine)
    with factory() as db:
        active = import_rule_catalog(db, 'master.xlsx', workbook_bytes(), activate=True)
        db.commit()
        candidate = import_rule_catalog(db, 'master-updated.xlsx', workbook_bytes(
            audit_check='Locate the approved Risk Register.'
        ))
        db.commit()
        active_version = db.get(ChecklistVersion, active['checklist_version_id'])
        candidate_version = db.get(ChecklistVersion, candidate['checklist_version_id'])
        assert active_version.is_active and active_version.status == 'ACTIVE'
        assert not candidate_version.is_active and candidate_version.status == 'VALIDATED'
        assert candidate['change_summary']['rules_changed'] == 1


def test_reimporting_the_same_workbook_is_idempotent():
    engine=create_engine('sqlite://'); Base.metadata.create_all(engine); factory=sessionmaker(bind=engine)
    content = workbook_bytes()
    with factory() as db:
        first = import_rule_catalog(db, 'master.xlsx', content)
        db.commit()
        repeated = import_rule_catalog(db, 'renamed-master.xlsx', content)
        db.commit()

        assert repeated['already_imported'] is True
        assert repeated['content_unchanged'] is True
        assert repeated['checklist_version_id'] == first['checklist_version_id']
        assert db.query(ChecklistVersion).count() == 1
        assert db.query(CmmiRule).count() == 1
        assert db.query(DocumentTypeRule).count() == 14


def test_document_key_updates_merge_only_into_an_approved_active_type_without_changing_rules():
    engine=create_engine('sqlite://'); Base.metadata.create_all(engine); factory=sessionmaker(bind=engine)
    with factory() as db:
        db.add_all([
            PracticeArea(code='EST', name='Estimating'),
            PracticeArea(code='PLAN', name='Planning'),
        ])
        active = import_rule_catalog(db, 'master.xlsx', workbook_bytes(), activate=True)
        db.commit()
        additions = document_additions_bytes([
            [1, 'WBS', 'EST/PLAN', 'Release work package, Approval milestone', 'Updated WBS'],
        ])

        result = import_document_catalogue_additions(db, 'document-additions.xlsx', additions)
        db.commit()

        assert result['checklist_version_id'] == active['checklist_version_id']
        assert len(result['source_checksum']) == 64
        assert result['document_types_inserted'] == 0
        assert result['document_types_updated'] == 1
        assert db.query(CmmiRule).count() == 1
        assert db.query(DocumentTypeRule).count() == 14
        wbs = db.scalar(select(DocumentTypeRule).where(DocumentTypeRule.document_type == 'WBS'))
        assert wbs.checklist_version_id == active['checklist_version_id']
        assert wbs.practice_areas == ['EST', 'PLAN']
        assert wbs.include_in_afr is True
        assert 'Release work package' in wbs.keywords
        assert 'Updated WBS' in wbs.expected_evidence

        repeated = import_document_catalogue_additions(db, 'document-additions.xlsx', additions)
        assert repeated['document_types_inserted'] == 0
        assert repeated['document_types_updated'] == 0
        assert repeated['document_types_unchanged'] == 1


def test_document_key_updates_reject_an_unapproved_document_type_without_mutation():
    engine=create_engine('sqlite://'); Base.metadata.create_all(engine); factory=sessionmaker(bind=engine)
    with factory() as db:
        db.add(PracticeArea(code='RSK', name='Risk'))
        active = import_rule_catalog(db, 'master.xlsx', workbook_bytes(), activate=True)
        db.commit()
        with pytest.raises(ValueError, match='not one of the approved 14'):
            import_document_catalogue_additions(db, 'document-additions.xlsx', document_additions_bytes([
                [1, 'Risk Register', 'RSK', 'Risk ID, Owner', 'Risk register'],
            ]))
        assert db.query(DocumentTypeRule).filter_by(checklist_version_id=active['checklist_version_id']).count() == 14


def test_document_additions_reject_corrupted_text_without_changing_the_active_catalogue():
    engine=create_engine('sqlite://'); Base.metadata.create_all(engine); factory=sessionmaker(bind=engine)
    with factory() as db:
        db.add(PracticeArea(code='RSK', name='Risk'))
        active = import_rule_catalog(db, 'master.xlsx', workbook_bytes(), activate=True)
        db.commit()
        corrupt = document_additions_bytes([
            [1, 'RCA for P1/P2', 'RSK', 'Owner, Status', 'Approved document'],
            [2, 'Incident Log � 2 Months', 'RSK', 'Incident ID, Status', 'Incident log'],
        ])

        try:
            import_document_catalogue_additions(db, 'document-additions.xlsx', corrupt)
        except ValueError as exc:
            assert 'invalid replacement character' in str(exc)
        else:
            raise AssertionError('corrupted catalogue text was accepted')
        assert db.query(DocumentTypeRule).filter_by(checklist_version_id=active['checklist_version_id']).count() == 14
        assert db.scalar(select(DocumentTypeRule).where(
            DocumentTypeRule.document_type == 'Valid New Document'
        )) is None


def test_document_key_updates_preserve_fixed_practice_area_mapping():
    engine=create_engine('sqlite://'); Base.metadata.create_all(engine); factory=sessionmaker(bind=engine)
    with factory() as db:
        db.add_all([PracticeArea(code='EST', name='Estimating'),
                    PracticeArea(code='PLAN', name='Planning'),
                    PracticeArea(code='MPM', name='Managing Performance')])
        active = import_rule_catalog(db, 'master.xlsx', workbook_bytes(), activate=True)
        db.commit()

        result = import_document_catalogue_additions(db, 'additions.xlsx', document_additions_bytes([
            [1, 'WBS', 'EST/PLAN', 'Variance, Status', 'Approved monitoring report'],
        ]))
        db.commit()

        document = db.scalar(select(DocumentTypeRule).where(
            DocumentTypeRule.checklist_version_id == active['checklist_version_id'],
            DocumentTypeRule.document_type == 'WBS',
        ))
        assert result['document_types_inserted'] == 0
        assert document.practice_areas == ['EST', 'PLAN']


def test_document_key_updates_accept_matching_pa_labels_but_reject_mapping_changes():
    engine=create_engine('sqlite://'); Base.metadata.create_all(engine); factory=sessionmaker(bind=engine)
    with factory() as db:
        db.add_all([
            PracticeArea(code='TS', name='Technical Solution'),
            PracticeArea(code='MPM', name='Managing Performance'),
        ])
        active = import_rule_catalog(db, 'master.xlsx', workbook_bytes(), activate=True)
        db.commit()
        accepted = import_document_catalogue_additions(db, 'additions.xlsx', document_additions_bytes([
            [1, 'Impact Analysis and TSD', 'Technical Solution (TS)', 'Architecture Decision', 'Approved design'],
        ]))
        db.commit()
        assert accepted['document_types_inserted'] == 0
        assert db.scalar(select(DocumentTypeRule).where(
            DocumentTypeRule.document_type == 'Impact Analysis and TSD'
        )).practice_areas == ['TS']

        try:
            import_document_catalogue_additions(db, 'additions.xlsx', document_additions_bytes([
                [1, 'Impact Analysis and TSD', 'TS/XYZ', 'Owner, Status', 'Approved summary'],
            ]))
        except ValueError as exc:
            assert 'unknown practice area(s): XYZ' in str(exc)
        else:
            raise AssertionError('an unknown code in a compact PA list was accepted')

        with pytest.raises(ValueError, match='fixed practice-area mapping'):
            import_document_catalogue_additions(db, 'additions.xlsx', document_additions_bytes([
                [1, 'Impact Analysis and TSD', 'MPM', 'Owner, Status', 'Approved summary'],
            ]))


def test_document_additions_reject_formulas_and_oversized_sheet_dimensions():
    engine=create_engine('sqlite://'); Base.metadata.create_all(engine); factory=sessionmaker(bind=engine)
    with factory() as db:
        db.add_all([
            PracticeArea(code='RSK', name='Risk'),
            ChecklistVersion(version='active', source='test', checksum='6' * 64,
                             status='ACTIVE', is_active=True),
        ])
        db.commit()
        formula = document_additions_bytes([
            [1, '=HYPERLINK("https://invalid.example")', 'RSK', 'Owner', 'Approved record'],
        ])
        with pytest.raises(ValueError, match='contains a formula'):
            import_document_catalogue_additions(db, 'additions.xlsx', formula)

        book = Workbook()
        sheet = book.active
        sheet.append(['S.No', 'Document / Evidence', 'Related Practice Area',
                      'What to Check / Key Points', 'Expected Evidence'])
        sheet.cell(row=2, column=201, value='dimension abuse')
        output = BytesIO(); book.save(output)
        with pytest.raises(ValueError, match='200-column import limit'):
            import_document_catalogue_additions(db, 'additions.xlsx', output.getvalue())
        assert db.query(DocumentTypeRule).count() == 0


def test_document_additions_require_exactly_one_active_version():
    engine=create_engine('sqlite://'); Base.metadata.create_all(engine); factory=sessionmaker(bind=engine)
    with factory() as db:
        db.add_all([
            PracticeArea(code='RSK', name='Risk'),
            ChecklistVersion(version='active-1', source='test', checksum='1' * 64,
                             status='ACTIVE', is_active=True),
            ChecklistVersion(version='active-2', source='test', checksum='2' * 64,
                             status='ACTIVE', is_active=True),
        ])
        db.commit()

        try:
            import_document_catalogue_additions(db, 'additions.xlsx', document_additions_bytes([
                [1, 'Incident Log', 'RSK', 'Incident ID, Status', 'Incident log'],
            ]))
        except ValueError as exc:
            assert 'exactly one active version' in str(exc)
        else:
            raise AssertionError('an ambiguous active ruleset was accepted')
        assert db.query(DocumentTypeRule).count() == 0


def test_document_additions_reject_a_high_compression_archive_before_mutation():
    engine=create_engine('sqlite://'); Base.metadata.create_all(engine); factory=sessionmaker(bind=engine)
    with factory() as db:
        db.add_all([
            PracticeArea(code='RSK', name='Risk'),
            ChecklistVersion(version='active', source='test', checksum='3' * 64,
                             status='ACTIVE', is_active=True),
        ])
        db.commit()
        payload = BytesIO()
        with zipfile.ZipFile(payload, 'w', compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('xl/worksheets/sheet1.xml', b'A' * (2 * 1024 * 1024))

        try:
            import_document_catalogue_additions(db, 'additions.xlsx', payload.getvalue())
        except ValueError as exc:
            assert 'corrupt, unsafe, or unreadable' in str(exc)
        else:
            raise AssertionError('a suspicious high-compression archive was accepted')
        assert db.query(DocumentTypeRule).count() == 0


def test_document_additions_reject_ambiguous_existing_normalized_names():
    engine=create_engine('sqlite://'); Base.metadata.create_all(engine); factory=sessionmaker(bind=engine)
    with factory() as db:
        db.add_all([
            PracticeArea(code='RSK', name='Risk'),
            ChecklistVersion(id=1, version='active', source='test', checksum='4' * 64,
                             status='ACTIVE', is_active=True),
            DocumentTypeRule(checklist_version_id=1, document_type='Risk Register',
                             practice_areas=['RSK'], keywords=['Risk ID'], aliases=[],
                             expected_evidence='Risk register'),
            DocumentTypeRule(checklist_version_id=1, document_type='Risk-Register',
                             practice_areas=['RSK'], keywords=['Owner'], aliases=[],
                             expected_evidence='Owned risks'),
        ])
        db.commit()

        try:
            import_document_catalogue_additions(db, 'additions.xlsx', document_additions_bytes([
                [1, 'Risk register', 'RSK', 'Status', 'Updated risk register'],
            ]))
        except ValueError as exc:
            assert 'ambiguous document names' in str(exc)
        else:
            raise AssertionError('ambiguous existing document names were silently merged')
        assert db.query(DocumentTypeRule).count() == 2


def test_invalid_catalog_does_not_create_or_activate_version():
    engine=create_engine('sqlite://'); Base.metadata.create_all(engine); factory=sessionmaker(bind=engine)
    with factory() as db:
        try: import_rule_catalog(db,'master.xlsx',workbook_bytes(duplicate=True))
        except ValueError as exc: assert 'Duplicate Rule ID' in str(exc)
        else: raise AssertionError('invalid workbook was accepted')
        assert db.scalar(select(ChecklistVersion.id)) is None


def test_rule_ids_are_canonicalized_and_duplicates_are_case_insensitive():
    engine=create_engine('sqlite://'); Base.metadata.create_all(engine); factory=sessionmaker(bind=engine)
    with factory() as db:
        try:
            import_rule_catalog(db, 'master.xlsx', workbook_bytes(
                duplicate=True, duplicate_rule_id='rsk-g1'
            ))
        except ValueError as exc:
            assert 'Duplicate Rule ID' in str(exc)
        else:
            raise AssertionError('a case-variant duplicate Rule ID was accepted')
        assert db.scalar(select(ChecklistVersion.id)) is None
