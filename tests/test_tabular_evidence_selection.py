from openpyxl import Workbook
from openpyxl.comments import Comment
import zipfile

from app.services.audit_engine.evidence_validation import build_tabular_context, classify_document
from app.services.audit_engine.header_matching import match_fields_exclusive
from app.services.audit_engine.risk_sla_validation import validate_risk_sla
from app.services.audit_engine.risk_sla_validation import FIELDS
from app.services.document_processing.extractors import extract_text


def test_risk_validator_selects_later_worksheet_and_real_header_row(tmp_path):
    source = tmp_path / 'customer-risk-register.xlsx'
    workbook = Workbook()
    revision = workbook.active
    revision.title = 'Revision History'
    revision.append(['Version', 'Date', 'Description'])
    revision.append([1, '2026-01-01', 'Initial release'])
    register = workbook.create_sheet('Risk Log')
    register.append(['Risk Log'])
    register.append(['ID#', 'Risk Level', 'Status', 'Date Identified', 'Date of Actual Closure', 'Remarks'])
    register.append(['RISK-1', 'High', 'Closed', '2026-01-01', '2026-01-03', ''])
    workbook.save(source)

    findings = validate_risk_sla(str(source))

    assert [item['rule_id'] for item in findings] == ['IRP-RISK-SLA', 'IRP-RISK-SLA-REASON']


def test_excel_extraction_includes_later_worksheets(tmp_path):
    source = tmp_path / 'traceability.xlsx'
    workbook = Workbook()
    workbook.active.title = 'Revision History'
    workbook.active.append(['Version', 'Date'])
    evidence = workbook.create_sheet('BRTM')
    evidence.append(['Requirement ID', 'System Test Cases'])
    evidence.append(['REQ-1', 'TC-1'])
    workbook.save(source)

    text = extract_text(str(source))

    assert '[Sheet: BRTM]' in text
    assert 'Requirement ID' in text


def test_excel_extraction_reads_merged_headers_hidden_sheets_comments_and_text_boxes(tmp_path):
    source = tmp_path / 'renamed-evidence.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Evidence'
    sheet.merge_cells('A1:D1')
    sheet['A1'] = 'Defect ID'
    sheet.append(['Severity', 'Priority', 'Status'])
    sheet['A2'].comment = Comment('Retest Closure', 'Auditor')
    hidden = workbook.create_sheet('Controlled Data')
    hidden.sheet_state = 'hidden'
    hidden.append(['Formula Evidence', '=CONCAT("approved", " fix")'])
    workbook.save(source)
    drawing_xml = '''<?xml version="1.0"?>
      <xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
                xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
        <xdr:twoCellAnchor><xdr:sp><xdr:txBody><a:p><a:r><a:t>Root Cause Corrective Fix</a:t></a:r></a:p></xdr:txBody></xdr:sp></xdr:twoCellAnchor>
      </xdr:wsDr>'''
    with zipfile.ZipFile(source, 'a') as package:
        package.writestr('xl/drawings/auditTextBox.xml', drawing_xml)

    text = extract_text(str(source))

    assert '[Sheet: Controlled Data]' in text
    assert 'Defect ID' in text
    assert 'CONCAT' in text
    assert 'Retest Closure' in text
    assert 'Root Cause Corrective Fix' in text
    result = classify_document(text, 'renamed-evidence.xlsx', [{
        'document_type': 'Testing Defects', 'practice_areas': ['VV'],
        'keywords': ['Defect ID', 'Severity', 'Priority', 'Root Cause', 'Status', 'Fix', 'Retest', 'Closure'],
        'include_in_afr': True,
    }])
    assert result['detected_type'] == 'Testing Defects'
    assert result['missed_keywords'] == []


def test_risk_validator_accepts_date_resolved_header(tmp_path):
    source = tmp_path / 'issue-log.csv'
    source.write_text(
        'Issue ID,Priority,Status,Date Raised,Date Resolved\n'
        'ISSUE-1,High,Closed,2026-01-01,2026-01-03\n', encoding='utf-8'
    )

    findings = validate_risk_sla(str(source))

    assert not [item for item in findings if item['rule_id'] == 'IRP-RISK-SLA-HEADER']


def test_later_header_is_found_after_large_leading_blank_area(tmp_path):
    source = tmp_path / 'spaced-risk-register.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Risk Log'
    sheet.cell(251, 1, 'Issue ID')
    sheet.cell(251, 2, 'Priority')
    sheet.cell(251, 3, 'Status')
    sheet.cell(251, 4, 'Date Raised')
    sheet.cell(251, 5, 'Date Resolved')
    sheet.append(['RISK-1', 'High', 'Closed', '2026-01-01', '2026-01-03'])
    workbook.save(source)

    findings = validate_risk_sla(str(source))

    assert not [item for item in findings if item['rule_id'] == 'IRP-RISK-SLA-HEADER']


def test_formula_text_is_preserved_for_document_classification(tmp_path):
    source = tmp_path / 'formula-evidence.xlsx'
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(['Requirement ID', 'Approval'])
    sheet.append(['REQ-1', '=CONCAT("approved", " baseline")'])
    workbook.save(source)

    text = extract_text(str(source))

    assert 'CONCAT' in text


def test_excel_extraction_preserves_zero_values(tmp_path):
    source = tmp_path / 'zero-evidence.xlsx'
    workbook = Workbook()
    workbook.active.append(['Variance %'])
    workbook.active.append([0])
    workbook.save(source)

    assert '0' in extract_text(str(source))


def test_tabular_control_context_preserves_formula_without_a_cached_result(tmp_path):
    source = tmp_path / 'formula-control.xlsx'
    workbook = Workbook()
    workbook.active.append(['Variance %'])
    workbook.active.append(['=E2/D2'])
    workbook.save(source)

    context = build_tabular_context(str(source), {'required_keywords': ['Variance %']})

    assert context['rows'] == [['=E2/D2']]


def test_fuzzy_matching_does_not_treat_resolution_as_a_closure_date():
    matched = match_fields_exclusive(['Resolution'], FIELDS)

    assert matched['closed'] is None


def test_detailed_srs_with_template_language_remains_implementation_evidence():
    result = classify_document(
        'Project scope stakeholders priority version requirements are approved. This template contains an example.',
        'Agriculture_SRS.docx',
        [{
            'document_type': 'BRD / SRS / Requirements Specification/RFP',
            'practice_areas': ['RDM'],
            'keywords': ['scope', 'stakeholders', 'priority', 'version', 'requirements', 'approved'],
        }],
    )

    assert result['detected_type'] == 'BRD / SRS / Requirements Specification/RFP'
    assert result['document_role'] == 'PROJECT_IMPLEMENTATION_EVIDENCE'
