from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'CMMI_Audit_Platform_Governance_Security_Assessment.docx'

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.7)
section.right_margin = Inches(0.7)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10)
styles['Normal'].paragraph_format.space_after = Pt(6)
for name, size in [('Title', 21), ('Heading 1', 15), ('Heading 2', 12)]:
    styles[name].font.name = 'Aptos Display'
    styles[name]._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos Display')
    styles[name]._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos Display')
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
styles['Heading 1'].paragraph_format.space_before = Pt(15)
styles['Heading 1'].paragraph_format.space_after = Pt(7)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def borders(cell, color='D9D9D9'):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = tcBorders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tcBorders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:color'), color)

def title(text, subtitle=False):
    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    if subtitle:
        r.font.name = 'Aptos'; r.font.size = Pt(10); r.font.bold = False; r.font.color.rgb = RGBColor(90,90,90)
    return p

def h(text):
    return doc.add_paragraph(text, style='Heading 1')

def para(text, bold=False, small=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if small: r.font.size = Pt(9)
    return p

def bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(text)
    p.paragraph_format.space_after = Pt(3)
    return p

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, header in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = header
        shade(cell, '305496')
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in cell.paragraphs[0].runs:
            r.font.name = 'Aptos'; r.font.size = Pt(8.5); r.font.bold = True; r.font.color.rgb = RGBColor(255,255,255)
        cell.width = Inches(widths[i]); borders(cell)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2: shade(cells[i], 'EAF2F8')
            cells[i].width = Inches(widths[i]); borders(cells[i])
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    r.font.name = 'Aptos'; r.font.size = Pt(8)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

title('CMMI Audit Platform Governance Security and Data Assessment')
title('Implementation based assessment of workflow access data protection update management and auditability', True)
para('Assessment scope: static review of the repository implementation and documented deployment design. This is not a penetration test, production configuration review, or independent operational evidence assessment.', small=True)

h('Overall conclusion')
para('The platform implements a defined audit workflow, authenticated role based access, email based account management, protected server side sessions, and persistent storage for audit evidence and derived records. It is suitable only as a shared internal workspace: screen access authorizes access to all customers and audit sessions visible on that screen, so it does not provide customer or organization isolation.')
para('The implementation has meaningful audit logging and deletion safeguards, but it should not be represented as maintaining a complete, tamper evident audit trail. Several business updates lack old and new values, there is no demonstrated log retention or export process, audit logs are not protected against a privileged database operator, and audit record access is not exposed as a controlled review function.')

h('1 Input Process Output workflow')
para('Yes. The application follows a defined Input to Process to Output workflow for a CMMI audit.', bold=True)
table(['Stage', 'Inputs', 'Processing and controls', 'Outputs'], [
['1 Setup', 'Customer name and contact email; project name; repository URL; audit session name and date; auditors and auditees.', 'Creates or reuses the customer, project and audit session. Pins the session to an active checklist version and seeds practice area tracking. Duplicate matching prevents an identical session from being created again.', 'A governed audit workspace with traceable project and session records.'],
['2 Evidence intake', 'Individual files, folders, or project ZIP archives. Supported office documents, PDFs, images and related formats are validated before persistence.', 'Checks size and extension; sanitizes filenames; validates ZIP traversal, member count, depth, compression ratio and symlinks; stores files outside static web content; calculates SHA-256 and identifies duplicates.', 'Evidence file records, provenance, storage paths and uploaded evidence retained for the session.'],
['3 Scan', 'Stored evidence, extracted text, active CMMI rules and document catalogue.', 'Extracts content, classifies documents, evaluates rules and key evidence, creates a scan job, records results and creates findings or review-required gaps.', 'Scan status, coverage status, findings, linked evidence and remediation context.'],
['4 Review and remediate', 'Findings, comments, remediation actions, due dates and status.', 'Authorized users review findings, add comments and track remediation actions.', 'Updated findings context and remediation history.'],
['5 Report', 'Audit session, scoped findings and report filters.', 'Generates AFR exports in supported formats and records generated-report metadata.', 'Downloadable audit findings reports plus generated report history.'],
], [0.85, 1.42, 2.55, 1.48])
para('Important boundary: the workflow is implemented as an application process, but it does not establish formal approval gates for each transition. A production procedure should define who may initiate a scan, approve findings, close remediation and release final reports.', small=True)

h('2 User roles and role based access control')
para('Yes, RBAC is implemented. Access is database backed through role to screen mappings. Each active screen has independent Read and Write grants; Write requires Read. The UI hides unauthorized menu entries, and API dependencies enforce the same screen-level permission checks. Administrators can create custom roles; a newly created role begins with no active screen access.')
table(['Role or control', 'Implemented access model', 'Assessment'], [
['Admin and Super Admin', 'Seeded roles have broad access, including user administration, role administration and rule catalogue management. Actual authorization relies on grants rather than role name.', 'Appropriate for controlled administrators. Separation of duties must be configured through custom roles and permission assignments.'],
['Auditor', 'May be granted operational screens, such as project setup, evidence scan and reports, according to the role screen matrix.', 'Supports auditor duties when least privilege grants are applied.'],
['Auditee', 'May be granted read access to relevant screens and any explicitly assigned write permissions.', 'The system supports the role, but it has no inherent per-project or per-customer boundary.'],
['Custom roles', 'An administrator can create roles and explicitly assign Read or Write for every active screen.', 'Flexible. Changes are logged, and safeguards prevent removal of the last active manager for user or role administration.'],
['Workspace scope', 'A user with screen access can access every customer, project, audit session, evidence item, finding and report exposed by that screen.', 'Material limitation. Do not use this release for segregating independent client organizations. Tenant or project membership controls are required.'],
], [1.15, 3.0, 2.15])
para('Authentication controls include Argon2id password hashes, expiring and revocable server-side sessions, HttpOnly session cookies, Secure cookies in production, SameSite=Lax, login rate limits, optional Turnstile verification, account lockout after configured failed attempts, and session revocation when a user is deactivated or changes a password.')

h('3 Personally identifiable information and project data')
para('Yes. The application stores PII and potentially sensitive project-related data. The exact content of uploaded evidence is controlled by the users who upload it, so it may itself contain further PII, confidential technical material or regulated data.', bold=True)
table(['Data category', 'Examples stored', 'Primary storage', 'Who can access'], [
['User account PII', 'Email address, display name, active status, role assignment, password hash, login failure and lockout state.', 'MySQL users, roles, sessions and reset tables.', 'Users and role administrators according to screen access; passwords are never stored in plaintext.'],
['Customer contact PII', 'Customer contact email and organization name.', 'MySQL customers table.', 'Any user with access to the shared Add Project workspace; not scoped to the creating user or one customer.'],
['Audit participant information', 'Auditor and auditee names, audit date and audit session name.', 'MySQL audit sessions and generated reports.', 'Users with relevant shared workspace or report screen access.'],
['Project information', 'Project name, repository URL, audit status, checklist version and provenance fields.', 'MySQL project and audit-session records.', 'Users with relevant screen access in the shared workspace.'],
['Evidence content', 'Documents, expanded archive members, filenames, folder paths, SHA-256, MIME type, size, extracted text and classifications.', 'Uploads volume; MySQL metadata and extracted text.', 'Users with Evidence Scan access. Deployment, database and volume administrators may also have infrastructure access.'],
['Audit outputs', 'Findings, comments, remediation, approvals, scan jobs, reports and history.', 'MySQL and reports volume.', 'Users with relevant shared screen access; report downloads require Reports access.'],
['Security and audit metadata', 'User IDs, actions, entity identifiers, timestamps, selected details, reset events and authorization denials.', 'MySQL audit_logs.', 'No dedicated audit-log viewer is shown in the reviewed UI/API. Database and infrastructure administrators can access it directly.'],
], [1.15, 2.0, 1.35, 1.8])
para('Files are deliberately not stored as database blobs. The database records a relative storage path and metadata; the physical file is retained on disk in the uploads volume. Database and volumes must therefore be backed up and access-controlled together.', small=True)

h('4 User email IDs and protection')
para('Yes. User email IDs are required account identifiers and are stored in the users table as unique, normalized values. Email IDs are also accepted for the first administrator, sign-in, password-reset requests and customer contact information. A failed sign-in and password-reset audit event may include the submitted email address in the audit-log detail.')
para('Implemented protections:', bold=True)
for x in [
    'Email input is validated and normalized; account emails are unique.',
    'Authentication passwords are protected with Argon2id hashing. Session and reset tokens are stored as SHA-256 digests rather than plaintext.',
    'Production configuration requires HTTPS and secure session cookies. The database runs on a private Docker network; the application port is intended to be exposed through an HTTPS reverse proxy.',
    'Password reset returns a non-enumerating response and rate limits reset requests. Login uses a dummy-hash check for unknown accounts to reduce account-enumeration timing signals.',
]: bullet(x)
para('Protection limitations and required operational controls: the reviewed schema does not encrypt email columns at rest, and audit logs can retain submitted email addresses. Production deployment must enforce encrypted disk or database storage, TLS for SMTP, least-privileged database accounts, restricted backup access, retention limits and a documented process for access, correction and deletion requests. Avoid placing PII in filenames and evidence content unless necessary for the audit.')

h('5 Technology stack and deployment responsibility')
table(['Layer', 'Technology', 'Operational responsibility'], [
['Web interface', 'NiceGUI, producing a browser interface using Quasar and Vue components.', 'Application owner maintains UI code and role configuration.'],
['Application and API', 'Python 3.12, FastAPI, Uvicorn, SQLAlchemy and Pydantic settings.', 'Application engineering team maintains code, tests and runtime configuration.'],
['Authentication and protection', 'argon2-cffi, SlowAPI, optional Cloudflare Turnstile and optional Microsoft OAuth configuration.', 'Security/application owner configures secrets, account controls, provider credentials and monitoring.'],
['Data', 'MySQL 8.4; persistent Docker volumes for MySQL, uploads, reports and templates.', 'Database/platform owner manages availability, backup, access privileges, encryption and recovery.'],
['Document processing', 'LibreOffice core, Tesseract OCR, PyMuPDF, pypdf, python-docx, python-pptx and related libraries.', 'Platform owner sets resource limits and malware-scanning controls for uploaded documents.'],
['Deployment', 'Docker and Docker Compose, behind an HTTPS reverse proxy on a Linux VPS.', 'The repository provides deployment instructions; it does not name an individual or team responsible for actual production deployment. Ownership must be assigned by the organization.'],
['Current state', 'Repository documentation states that the tool is currently run locally for development/testing; the containerized production path is prepared.', 'There is no evidence in the repository of a live production deployment or designated production operator.'],
], [1.15, 2.85, 2.3])
para('Production design requires a non-root runtime database account, a separate DDL-capable migration account, a unique 32+ character secret key, an HTTPS public URL, production migration disabled at application startup, and volumes backed up together.')

h('6 Security patches and updates')
para('The repository defines a security assurance pipeline, but it does not yet define a complete operational patch-management policy or release approval process.', bold=True)
table(['Control area', 'Implemented evidence', 'Gap or action required'], [
['Dependency checks', 'GitHub Actions installs pip-audit and runs pip check plus pip-audit against requirements.txt.', 'Dependency versions are compatibility ranges, not a fully hash-pinned production lockfile. Generate, review and maintain a pinned lockfile and SBOM per release.'],
['Static and secret scanning', 'CI runs Bandit against application code and Gitleaks against repository history.', 'Define triage, remediation SLA, ownership and exception approval for findings.'],
['Container checks', 'CI builds the Docker image and scans it with Trivy for High and Critical vulnerabilities.', 'Document base image update cadence, accepted exceptions and re-scan requirements.'],
['Update testing', 'Repository has pytest coverage and deployment guidance.', 'The recorded workspace result says tests and Docker build were not executable there because Python and Docker were unavailable. Before release, run tests, security workflow and a staging deployment with evidence of results.'],
['Deployment of changes', 'Docker Compose supports a build-and-restart path with persistent volumes; schema changes use a separate migration job.', 'Define formal stages: dependency review, test, security scan, staging validation, backup, approval, production release, smoke test and rollback.'],
['Emergency security patch', 'No dedicated emergency patch procedure is documented.', 'Establish severity-based emergency change authority, expedited test scope, communications, monitoring and post-implementation review.'],
], [1.35, 2.85, 2.1])

h('7 Audit trail historical records and deletion prevention')
para('Partial yes. The application has an audit_logs table and records many security and administrative events. Each event records a user ID where known, action, entity type, entity ID, JSON detail and a UTC timestamp. The application code describes this log as application-write-only: no reviewed GUI or API mutation or deletion path is exposed for audit-log records.')
table(['Area', 'Implemented control', 'Assessment'], [
['Events captured', 'Examples include sign-in success/failure, first-administrator registration, password reset/change, user lifecycle, role changes, screen permission changes, rule catalogue actions, evidence upload/replacement, scan start/complete/failure, report generation and authorization denial.', 'Meaningful event coverage.'],
['Change details', 'The logger can store structured changes. The remediation update records field names as changed. Many actions record only an action and entity identifier.', 'Not complete reconstruction: old and new values are generally absent, and some business mutations have no comprehensive change payload.'],
['Audit-log integrity', 'Application UI/API has no route to update or delete audit-log entries. Runtime database account is designed to lack DDL privileges.', 'Not tamper-evident or immutable. A privileged database or infrastructure administrator can alter rows unless database-level controls and external immutable retention are added.'],
['Business-record deletion', 'Project/session deletion is blocked when evidence, scans, findings, approvals, reports or incident-validation records exist. Users are normally deactivated and sessions revoked rather than removed. Unreadable evidence is retained and marked replaced.', 'Good application-level preservation controls, but some models use cascading foreign keys and database-level deletion or migrations can still alter history.'],
['Historical retention', 'Timestamped operational tables retain created and updated times; checklist version is pinned to audit sessions; generated reports and source evidence metadata persist.', 'No documented retention schedule, legal hold, backup retention, restore testing, audit-log export, or independent audit-log review process.'],
['Access to audit trail', 'Audit rows are stored in MySQL.', 'A controlled audit-log inquiry/export interface and auditor role are not evident in the reviewed UI/API.'],
], [1.25, 2.95, 2.1])
para('Conclusion for this question: do not state that the platform maintains a complete audit trail of every user activity and every data change. It maintains a useful application-level log and preserves key audit artifacts, but it needs coverage validation, before/after values, immutable external retention, controlled review access and formal retention/deletion governance to meet a stronger auditability claim.', bold=True)

h('Priority actions before a production assurance claim')
for x in [
    'Assign named owners for platform deployment, database operations, security patching, incident response, privacy requests and audit-log review.',
    'Introduce organization/project membership and data scoping if customers must be isolated from each other.',
    'Protect audit history with database write restrictions, immutable or WORM off-platform export, integrity checks, retention schedules, backup/restore testing and a reviewed audit-log export function.',
    'Capture old and new values, actor, source IP/session context and outcome for all sensitive create/update/delete events; verify coverage with automated tests.',
    'Adopt a formal patch and release process with pinned dependencies, vulnerability remediation SLAs, staging validation and documented production approvals.',
    'Apply encryption at rest for MySQL, files and backups; restrict evidence and audit-log access; publish data retention and PII handling rules.',
]: bullet(x)

h('Repository evidence reviewed')
para('README.md; docs/architecture.md; docs/security.md; docs/deployment.md; docs/configuration.md; docs/test_results.md; Dockerfile; docker-compose.yml; requirements.txt; .github/workflows/security.yml; app/db/models.py; app/core/audit_log.py; app/core/authorization.py; app/core/security.py; app/core/config.py; app/core/screen_registry.py; app/core/middleware.py; app/services/auth/service.py; app/services/user_administration.py; app/services/role_administration.py; app/services/evidence_ingestion.py; app/services/document_processing/file_security.py; and the reviewed FastAPI endpoint modules.', small=True)

for sec in doc.sections:
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run('CMMI Audit Platform Governance Security and Data Assessment')
    run.font.name = 'Aptos'; run.font.size = Pt(8); run.font.color.rgb = RGBColor(128,128,128)

doc.save(OUT)
print(OUT)
